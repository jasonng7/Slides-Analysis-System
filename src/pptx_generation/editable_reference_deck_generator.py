from __future__ import annotations

from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from pptx.util import Inches

from config import OUTPUT_DIRECTORIES
from pptx_generation.generation_models import GeneratedDeckInput, GeneratedDeckMetadata
from pptx_generation.notes_builder import build_slide_notes
from pptx_generation.placeholder_deck_generator import (
    _deck_title,
    _extract_patterns,
    _pattern_matching_method,
    _referenced_template_slides,
    _user_goal,
    load_recommendation_json,
    normalize_recommendation_to_slide_plans,
)
from pptx_generation.pptx_repository import (
    default_output_path,
    ensure_generated_deck_dirs,
    save_generation_metadata,
    save_speaker_notes_markdown,
)
from pptx_generation.slide_layout_builder import (
    SLIDE_H,
    SLIDE_W,
    add_agenda_slide,
    add_appendix_slide,
    add_cover_slide,
)
from pptx_generation.source_slide_copier import (
    SlideCloneError,
    add_image_fallback_slide,
    clone_reference_slide_into,
)
from recommendation.recommend_from_content import recommend_from_content


def generate_editable_reference_deck_from_recommendation(
    recommendation_path: str | Path,
    output_path: str | Path | None = None,
) -> dict[str, Any]:
    recommendation = load_recommendation_json(recommendation_path)
    return _generate_editable_reference_deck(
        recommendation=recommendation,
        source_recommendation_path=recommendation.get("_source_recommendation_path") or str(recommendation_path),
        output_path=output_path,
    )


def generate_editable_reference_deck_from_text(
    text: str,
    goal: str | None,
    mode: str = "deck",
    output_path: str | Path | None = None,
) -> dict[str, Any]:
    recommendation = recommend_from_content(text=text, user_goal=goal, mode=mode, prompt_output=False)
    return _generate_editable_reference_deck(
        recommendation=recommendation,
        source_recommendation_path=recommendation.get("recommendation_json_path"),
        output_path=output_path,
    )


def generate_editable_reference_deck_from_file(
    file_path: str | Path,
    goal: str | None,
    mode: str = "deck",
    output_path: str | Path | None = None,
) -> dict[str, Any]:
    recommendation = recommend_from_content(
        file_path=str(file_path),
        user_goal=goal,
        mode=mode,
        prompt_output=False,
    )
    return _generate_editable_reference_deck(
        recommendation=recommendation,
        source_recommendation_path=recommendation.get("recommendation_json_path"),
        output_path=output_path,
    )


def _generate_editable_reference_deck(
    *,
    recommendation: dict[str, Any],
    source_recommendation_path: str | None,
    output_path: str | Path | None,
) -> dict[str, Any]:
    try:
        from pptx import Presentation
    except ImportError as exc:
        raise ImportError(
            "python-pptx is required for editable reference deck generation. "
            "Install it with: .venv/bin/python -m pip install python-pptx"
        ) from exc

    ensure_generated_deck_dirs()
    slide_plans = normalize_recommendation_to_slide_plans(recommendation)
    if not slide_plans:
        raise ValueError("No slide plans found in recommendation JSON.")

    timestamp = datetime.now(timezone.utc).strftime("%Y%m%d_%H%M%S")
    title = _deck_title(recommendation)
    target_path = Path(output_path) if output_path else default_output_path(f"{title}_editable_reference", timestamp)
    if not target_path.is_absolute():
        target_path = Path.cwd() / target_path
    target_path.parent.mkdir(parents=True, exist_ok=True)

    deck_input = GeneratedDeckInput(
        title=title,
        user_goal=_user_goal(recommendation),
        mode=recommendation.get("recommended_output_mode", "deck"),
        source_recommendation_path=source_recommendation_path,
        matching_method=recommendation.get("matching_method") or _pattern_matching_method(_extract_patterns(recommendation)),
        slide_patterns=_extract_patterns(recommendation),
        recommendation=recommendation,
    ).as_dict()

    prs = Presentation()
    prs.slide_width = Inches(SLIDE_W)
    prs.slide_height = Inches(SLIDE_H)
    add_cover_slide(prs, {**deck_input, "title": f"{title}\nEditable reference-slide draft"})
    add_agenda_slide(prs, slide_plans)

    cloned_count = 0
    fallback_count = 0
    clone_failures: list[dict[str, Any]] = []
    cloned_references: list[dict[str, Any]] = []
    presentation_cache: dict[str, Any] = {}

    for slide_plan in slide_plans:
        try:
            _, reference = clone_reference_slide_into(prs, slide_plan, presentation_cache)
            cloned_count += 1
            cloned_references.append(
                {
                    "deck_name": reference.get("deck_name"),
                    "slide_number": reference.get("slide_number"),
                    "slide_title": reference.get("slide_title"),
                    "source_file": reference.get("source_file"),
                }
            )
        except Exception as exc:
            fallback_count += 1
            clone_failures.append(
                {
                    "slide_title": slide_plan.get("slide_title"),
                    "matched_template": slide_plan.get("matched_template"),
                    "error": str(exc),
                }
            )
            add_image_fallback_slide(prs, slide_plan)

    add_appendix_slide(prs, deck_input, slide_plans)
    prs.save(target_path)

    notes_path = save_speaker_notes_markdown(
        [
            {
                "slide_number": str(index),
                "slide_title": slide_plan.get("slide_title", f"Slide {index}"),
                "notes": build_slide_notes(slide_plan),
            }
            for index, slide_plan in enumerate(slide_plans, start=1)
        ],
        target_path,
    )

    warnings = [
        "Stage 8B duplicates source PPTX slide shapes where possible; complex charts, groups, media, or theme relationships may still require manual QA.",
    ]
    if fallback_count:
        warnings.append(f"{fallback_count} slide(s) fell back to Stage 8A image-background mode.")

    metadata = GeneratedDeckMetadata(
        generated_at=datetime.now(timezone.utc).isoformat(),
        output_path=str(target_path),
        source_recommendation_path=source_recommendation_path,
        user_goal=deck_input.get("user_goal", ""),
        mode=deck_input.get("mode", "deck"),
        number_of_slides=len(prs.slides),
        number_of_content_slides=len(slide_plans),
        matching_method=deck_input.get("matching_method", "metadata_match"),
        referenced_template_slides=_referenced_template_slides(slide_plans),
        warnings=warnings,
        generator_version="stage8b_editable_reference_slide_clone_v1",
    ).as_dict()
    metadata["speaker_notes_path"] = str(notes_path)
    metadata["reference_background_mode"] = "pptx_shape_clone_with_image_fallback"
    metadata["content_slides_cloned_from_source_pptx"] = cloned_count
    metadata["content_slides_with_image_fallback"] = fallback_count
    metadata["cloned_reference_slides"] = cloned_references
    metadata["clone_failures"] = clone_failures

    metadata_path = save_generation_metadata(metadata, target_path)
    metadata["metadata_path"] = str(metadata_path)
    log_path = OUTPUT_DIRECTORIES["generated_deck_logs"] / f"{target_path.stem}.log"
    log_path.write_text(
        "\n".join(
            [
                f"Generated editable reference deck: {target_path}",
                f"Slides: {len(prs.slides)}",
                f"Cloned slides: {cloned_count}/{len(slide_plans)}",
                f"Fallback slides: {fallback_count}",
                f"Metadata: {metadata_path}",
            ]
        )
        + "\n",
        encoding="utf-8",
    )
    metadata["log_path"] = str(log_path)
    return metadata
