from __future__ import annotations

from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from config import OUTPUT_DIRECTORIES, SUPPORTED_CONTENT_EXTENSIONS
from utils.file_utils import ensure_project_directories, resolve_input_file, slugify


class UnsupportedContentTypeError(ValueError):
    pass


def extract_content(file_path: str | None = None, text: str | None = None) -> dict[str, Any]:
    ensure_project_directories()

    if not file_path and text is None:
        raise ValueError("Provide either file_path or text.")
    if file_path and text is not None:
        raise ValueError("Provide only one of file_path or text.")

    timestamp = _timestamp()

    if text is not None:
        source_name = "pasted_text"
        source_type = "text"
        raw_text = text.strip()
        source_file = None
        metadata = {"input_method": "pasted_text", "character_count": len(raw_text)}
    else:
        path = resolve_input_file(file_path or "")
        if not path.exists():
            raise FileNotFoundError(f"Content file not found: {path}")
        extension = path.suffix.lower()
        if extension not in SUPPORTED_CONTENT_EXTENSIONS:
            supported = ", ".join(sorted(SUPPORTED_CONTENT_EXTENSIONS))
            raise UnsupportedContentTypeError(
                f"Unsupported content type '{extension}'. Supported: {supported}"
            )

        source_name = path.stem
        source_type = extension.lstrip(".")
        source_file = str(path)
        raw_text, metadata = _extract_from_file(path, source_type)

    raw_text_path = _write_raw_text(source_name, timestamp, raw_text)
    return {
        "source_name": source_name,
        "source_type": source_type,
        "source_file": source_file,
        "raw_text": raw_text,
        "raw_text_path": str(raw_text_path),
        "timestamp": timestamp,
        "metadata": metadata,
    }


def _extract_from_file(path: Path, source_type: str) -> tuple[str, dict[str, Any]]:
    if source_type == "pdf":
        return _extract_pdf(path)
    if source_type == "pptx":
        return _extract_pptx(path)
    if source_type == "docx":
        return _extract_docx(path)
    if source_type in {"txt", "md"}:
        text = path.read_text(encoding="utf-8", errors="replace")
        return text, {"character_count": len(text)}
    if source_type == "csv":
        return _extract_csv(path)
    if source_type == "xlsx":
        return _extract_xlsx_placeholder(path)
    raise UnsupportedContentTypeError(f"Unsupported content type: {source_type}")


def _extract_pdf(path: Path) -> tuple[str, dict[str, Any]]:
    import fitz

    fitz.TOOLS.mupdf_display_errors(False)
    fitz.TOOLS.mupdf_display_warnings(False)

    page_text: list[str] = []
    with fitz.open(path) as document:
        for index, page in enumerate(document, start=1):
            text = page.get_text("text").strip()
            page_text.append(f"## Page {index}\n{text}" if text else f"## Page {index}\n")
        metadata = {"page_count": document.page_count}

    return "\n\n".join(page_text).strip(), metadata


def _extract_pptx(path: Path) -> tuple[str, dict[str, Any]]:
    from pptx import Presentation

    try:
        presentation = Presentation(path)
        slide_texts: list[str] = []

        for slide_index, slide in enumerate(presentation.slides, start=1):
            chunks: list[str] = []
            for shape in slide.shapes:
                chunks.extend(_shape_text(shape))
            slide_text = "\n".join(dict.fromkeys(chunks)).strip()
            slide_texts.append(f"## Slide {slide_index}\n{slide_text}")

        return "\n\n".join(slide_texts).strip(), {"slide_count": len(presentation.slides)}
    except Exception:
        return _extract_pptx_xml_fallback(path)


def _shape_text(shape) -> list[str]:
    chunks: list[str] = []
    if hasattr(shape, "text") and shape.text:
        chunks.append(shape.text.strip())
    if getattr(shape, "has_table", False):
        for row in shape.table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                chunks.append(" | ".join(cells))
    if getattr(shape, "shape_type", None) == 6 and hasattr(shape, "shapes"):
        for nested_shape in shape.shapes:
            chunks.extend(_shape_text(nested_shape))
    return [chunk for chunk in chunks if chunk]


def _extract_pptx_xml_fallback(path: Path) -> tuple[str, dict[str, Any]]:
    """Extract PPTX text directly from slide XML when python-pptx cannot open it."""
    import re
    import zipfile
    import xml.etree.ElementTree as ET

    slide_texts: list[str] = []
    namespaces = {
        "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    }
    with zipfile.ZipFile(path) as archive:
        slide_names = sorted(
            [name for name in archive.namelist() if re.match(r"ppt/slides/slide\d+\.xml$", name)],
            key=lambda name: int(re.search(r"slide(\d+)\.xml$", name).group(1)),  # type: ignore[union-attr]
        )
        for slide_index, name in enumerate(slide_names, start=1):
            root = ET.fromstring(archive.read(name))
            chunks = [
                node.text.strip()
                for node in root.findall(".//a:t", namespaces)
                if node.text and node.text.strip()
            ]
            slide_texts.append(f"## Slide {slide_index}\n" + "\n".join(dict.fromkeys(chunks)))
    return "\n\n".join(slide_texts).strip(), {
        "slide_count": len(slide_texts),
        "extraction_fallback": "pptx_xml",
    }


def _extract_docx(path: Path) -> tuple[str, dict[str, Any]]:
    from docx import Document

    document = Document(path)
    chunks: list[str] = []

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text:
            chunks.append(text)

    for table_index, table in enumerate(document.tables, start=1):
        chunks.append(f"## Table {table_index}")
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                chunks.append(" | ".join(cells))

    return "\n\n".join(chunks).strip(), {
        "paragraph_count": len(document.paragraphs),
        "table_count": len(document.tables),
    }


def _extract_csv(path: Path) -> tuple[str, dict[str, Any]]:
    import pandas as pd

    dataframe = pd.read_csv(path)
    sections = [
        f"# CSV summary for {path.name}",
        f"Rows: {len(dataframe)}",
        f"Columns: {len(dataframe.columns)}",
        "Column names: " + ", ".join(str(column) for column in dataframe.columns),
        "## Sample rows",
        "```text\n" + dataframe.head(10).to_string(index=False) + "\n```",
    ]

    numeric_description = dataframe.describe(include="all").fillna("").to_string()
    sections.extend(["## Descriptive statistics", numeric_description])

    return "\n\n".join(sections), {
        "row_count": len(dataframe),
        "column_count": len(dataframe.columns),
        "columns": [str(column) for column in dataframe.columns],
    }


def _extract_xlsx_placeholder(path: Path) -> tuple[str, dict[str, Any]]:
    # TODO: Later extract each sheet, columns, sample rows, shape, formulas,
    # and possible chart recommendations from Excel workbooks.
    raise NotImplementedError("XLSX support will be implemented in a later stage.")


def _write_raw_text(source_name: str, timestamp: str, raw_text: str) -> Path:
    output_dir = OUTPUT_DIRECTORIES["content_raw_text"]
    output_dir.mkdir(parents=True, exist_ok=True)
    path = output_dir / f"{slugify(source_name)}_{timestamp}.txt"
    path.write_text(raw_text + ("\n" if raw_text else ""), encoding="utf-8")
    return path


def _timestamp() -> str:
    return datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%SZ")
