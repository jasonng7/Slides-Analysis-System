from __future__ import annotations

from typing import Any

from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

from pptx_generation.notes_builder import build_slide_notes


SLIDE_W = 13.333
SLIDE_H = 7.5
MARGIN_X = 0.45
DARK = RGBColor(30, 36, 45)
MID = RGBColor(92, 103, 115)
LIGHT = RGBColor(242, 244, 247)
BORDER = RGBColor(165, 174, 185)
ACCENT = RGBColor(10, 64, 116)
PALE_BLUE = RGBColor(234, 241, 248)
PALE_PEACH = RGBColor(250, 238, 228)


def add_cover_slide(prs, deck_input: dict[str, Any]) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_background(slide)
    title = deck_input.get("title") or "Generated Placeholder Deck"
    _add_textbox(slide, MARGIN_X, 1.0, 11.6, 1.2, title, 34, bold=True, color=DARK)
    _add_textbox(
        slide,
        MARGIN_X,
        2.35,
        7.8,
        0.45,
        "Generated placeholder deck",
        18,
        bold=True,
        color=ACCENT,
    )
    _add_textbox(
        slide,
        MARGIN_X,
        3.05,
        8.5,
        0.9,
        "Editable first draft for consultant review. Placeholder content, layouts, and build notes are intended for manual refinement.",
        15,
        color=MID,
    )
    details = [
        f"Mode: {deck_input.get('mode') or 'deck'}",
        f"Matching method: {deck_input.get('matching_method') or 'metadata_match'}",
        f"Source recommendation: {deck_input.get('source_recommendation_path') or 'Generated directly from content'}",
    ]
    _add_placeholder_box(slide, MARGIN_X, 4.65, 7.7, 1.3, "\n".join(details), fill=PALE_BLUE)
    add_footer(slide, "Stage 7 placeholder draft", None)


def add_agenda_slide(prs, slide_plans: list[dict[str, Any]]) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_background(slide)
    add_title(slide, "Agenda")
    headers = ["#", "Slide title", "Type", "Purpose"]
    rows = [
        [
            str(index),
            plan.get("slide_title") or f"Slide {index}",
            plan.get("slide_type") or "other",
            plan.get("slide_objective") or plan.get("storyline") or "Draft slide skeleton",
        ]
        for index, plan in enumerate(slide_plans, start=1)
    ]
    add_table_placeholder(slide, 0.65, 1.25, 12.0, 5.45, headers, rows[:12])
    if len(rows) > 12:
        _add_textbox(
            slide,
            0.65,
            6.78,
            7.0,
            0.25,
            f"+ {len(rows) - 12} additional slide(s) included after this agenda page.",
            9,
            color=MID,
        )
    add_footer(slide, "Generated agenda", None)


def add_content_slide(prs, slide_plan: dict[str, Any], slide_number: int) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_background(slide)
    add_title(slide, slide_plan.get("slide_title") or f"Placeholder slide {slide_number}")

    layout_family = select_layout_family(slide_plan)
    if layout_family == "executive_summary":
        _add_executive_summary_layout(slide, slide_plan)
    elif layout_family == "competitor_benchmark":
        _add_competitor_layout(slide, slide_plan)
    elif layout_family == "market_landscape":
        _add_market_landscape_layout(slide, slide_plan)
    elif layout_family == "market_sizing":
        _add_market_sizing_layout(slide, slide_plan)
    elif layout_family == "strategic_options":
        add_option_comparison_placeholder(slide, slide_plan)
    elif layout_family == "roadmap":
        add_timeline_placeholder(slide, slide_plan)
    elif layout_family == "risk":
        add_risk_matrix_placeholder(slide, slide_plan)
    else:
        _add_generic_layout(slide, slide_plan)

    _add_build_notes_box(slide, slide_plan)
    add_footer(slide, _footer_text(slide_plan), slide_plan)


def add_appendix_slide(prs, deck_input: dict[str, Any], slide_plans: list[dict[str, Any]]) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_background(slide)
    add_title(slide, "Appendix: matched template patterns")

    patterns = deck_input.get("slide_patterns") or []
    lines = [
        f"Matching method used: {deck_input.get('matching_method') or 'metadata_match'}",
        "Limitations: this is a placeholder draft, not final design automation.",
        "Next step: replace placeholders with verified evidence, charts, and consultant-edited copy.",
    ]
    _add_placeholder_box(slide, 0.65, 1.1, 12.0, 1.0, "\n".join(lines), fill=PALE_PEACH)

    rows = []
    for pattern in patterns[:8]:
        score = pattern.get("similarity_score") or pattern.get("match_score") or ""
        rows.append(
            [
                pattern.get("deck_name") or "",
                str(pattern.get("slide_number") or ""),
                pattern.get("slide_title") or "",
                pattern.get("slide_type") or "",
                str(score),
            ]
        )
    if not rows:
        rows = [["No matched templates available", "", "", "", ""]]
    add_table_placeholder(
        slide,
        0.65,
        2.45,
        12.0,
        3.9,
        ["Deck", "Slide", "Title", "Type", "Score"],
        rows,
    )
    add_footer(slide, "Template reference appendix", None)


def add_title(slide, text: str) -> None:
    _add_textbox(slide, MARGIN_X, 0.28, 12.3, 0.65, text, 21, bold=True, color=DARK)
    line = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.RECTANGLE,
        Inches(MARGIN_X),
        Inches(0.95),
        Inches(12.35),
        Inches(0.02),
    )
    line.fill.solid()
    line.fill.fore_color.rgb = ACCENT
    line.line.fill.background()


def add_footer(slide, text: str, slide_plan: dict[str, Any] | None) -> None:
    footer = text
    if slide_plan:
        matched = slide_plan.get("matched_template") or {}
        bits = [
            matched.get("deck_name"),
            f"slide {matched.get('slide_number')}" if matched.get("slide_number") else None,
            slide_plan.get("slide_type"),
            slide_plan.get("matching_method"),
        ]
        score = matched.get("similarity_score") or matched.get("match_score")
        if score is not None:
            bits.append(f"score {score}")
        footer = " | ".join(str(bit) for bit in bits if bit)
    _add_textbox(slide, MARGIN_X, 7.08, 12.2, 0.22, footer, 7, color=MID)


def add_placeholder_box(
    slide,
    x: float,
    y: float,
    w: float,
    h: float,
    text: str,
    *,
    fill=LIGHT,
    border=BORDER,
    font_size: int = 10,
) -> Any:
    return _add_placeholder_box(slide, x, y, w, h, text, fill=fill, border=border, font_size=font_size)


def add_table_placeholder(
    slide,
    x: float,
    y: float,
    w: float,
    h: float,
    headers: list[str],
    rows: list[list[str]],
) -> Any:
    row_count = max(2, min(len(rows) + 1, 14))
    col_count = max(1, len(headers))
    table_shape = slide.shapes.add_table(row_count, col_count, Inches(x), Inches(y), Inches(w), Inches(h))
    table = table_shape.table
    for col in range(col_count):
        table.cell(0, col).text = headers[col]
        _format_cell(table.cell(0, col), fill=ACCENT, color=RGBColor(255, 255, 255), bold=True)
    for row_index, row in enumerate(rows[: row_count - 1], start=1):
        for col_index in range(col_count):
            table.cell(row_index, col_index).text = row[col_index] if col_index < len(row) else ""
            _format_cell(table.cell(row_index, col_index), fill=RGBColor(255, 255, 255), color=DARK)
    return table_shape


def add_chart_placeholder(slide, x: float, y: float, w: float, h: float, label: str = "Chart placeholder") -> Any:
    box = _add_placeholder_box(slide, x, y, w, h, label, fill=RGBColor(248, 249, 251), font_size=12)
    for i, height in enumerate([0.55, 0.95, 1.3, 0.8]):
        bar = slide.shapes.add_shape(
            MSO_AUTO_SHAPE_TYPE.RECTANGLE,
            Inches(x + 0.65 + i * 0.65),
            Inches(y + h - 0.45 - height),
            Inches(0.32),
            Inches(height),
        )
        bar.fill.solid()
        bar.fill.fore_color.rgb = ACCENT if i % 2 == 0 else RGBColor(103, 136, 166)
        bar.line.fill.background()
    return box


def add_timeline_placeholder(slide, slide_plan: dict[str, Any]) -> None:
    phases = _content_or_defaults(slide_plan, ["Phase 1", "Phase 2", "Phase 3", "Phase 4"])
    _add_textbox(slide, 0.65, 1.22, 11.5, 0.35, "Timeline / implementation roadmap placeholder", 12, bold=True, color=DARK)
    y = 2.1
    for index, phase in enumerate(phases[:4]):
        x = 0.8 + index * 3.0
        _add_placeholder_box(slide, x, y, 2.55, 1.2, f"{phase}\nMilestones\nOwner / timing", fill=PALE_BLUE)
        arrow = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RIGHT_ARROW, Inches(x + 2.25), Inches(y + 0.38), Inches(0.65), Inches(0.35))
        arrow.fill.solid()
        arrow.fill.fore_color.rgb = RGBColor(190, 198, 208)
        arrow.line.fill.background()
    _add_placeholder_box(slide, 0.8, 4.25, 7.6, 0.9, "Owners / dependencies / decision gates", fill=LIGHT)


def add_option_comparison_placeholder(slide, slide_plan: dict[str, Any]) -> None:
    options = ["Build", "Buy", "Partner"]
    for index, option in enumerate(options):
        _add_placeholder_box(
            slide,
            0.7 + index * 4.05,
            1.35,
            3.75,
            3.2,
            f"{option}\n\nValue proposition\nInvestment required\nSpeed to market\nControl / risk",
            fill=RGBColor(255, 255, 255),
        )
    _add_placeholder_box(slide, 0.7, 4.85, 7.9, 0.75, "Evaluation criteria row: strategic fit | economics | execution risk | regulatory risk", fill=PALE_BLUE)
    _add_placeholder_box(slide, 8.9, 4.85, 3.55, 0.75, "Recommendation / preferred option", fill=PALE_PEACH)


def add_risk_matrix_placeholder(slide, slide_plan: dict[str, Any]) -> None:
    add_table_placeholder(
        slide,
        0.7,
        1.35,
        7.95,
        3.8,
        ["Risk / issue", "Severity", "Mitigation", "Owner"],
        [
            ["Regulatory", "High / med / low", "Mitigation action", "Owner"],
            ["Execution", "High / med / low", "Mitigation action", "Owner"],
            ["Commercial", "High / med / low", "Mitigation action", "Owner"],
            ["Operational", "High / med / low", "Mitigation action", "Owner"],
        ],
    )
    _add_placeholder_box(slide, 8.95, 1.35, 3.45, 1.5, "Priority / severity map", fill=PALE_BLUE)
    _add_placeholder_box(slide, 8.95, 3.1, 3.45, 2.05, "Decision implication and conditions to proceed", fill=PALE_PEACH)


def select_layout_family(slide_plan: dict[str, Any]) -> str:
    slide_type = str(slide_plan.get("slide_type") or "").lower()
    chart_type = str(slide_plan.get("chart_type") or "").lower()
    layout = str(slide_plan.get("layout_pattern") or "").lower()
    haystack = " ".join([slide_type, chart_type, layout, str(slide_plan.get("slide_title") or "").lower()])
    if slide_type in {"executive_summary", "board_decision_paper", "recommendation_slide"}:
        return "executive_summary"
    if "competitor" in haystack or "benchmark" in haystack or "peer" in haystack:
        return "competitor_benchmark"
    if "landscape" in haystack or "market_attractiveness" in haystack:
        return "market_landscape"
    if slide_type in {"market_sizing", "financial_summary", "valuation_comparison"} or chart_type in {"bar_chart", "waterfall_chart", "line_chart"}:
        return "market_sizing"
    if "option" in haystack or "build" in haystack or "buy" in haystack or "partner" in haystack or slide_type in {"options_analysis", "prioritization_matrix", "growth_options_analysis"}:
        return "strategic_options"
    if slide_type in {"roadmap", "implementation_plan", "gantt_timeline"} or "timeline" in haystack:
        return "roadmap"
    if slide_type == "risk_assessment" or "risk" in haystack or "mitigation" in haystack:
        return "risk"
    return "generic"


def _add_executive_summary_layout(slide, slide_plan: dict[str, Any]) -> None:
    items = _content_or_defaults(slide_plan, ["Situation", "Complication", "Resolution"])
    for index, item in enumerate(items[:5]):
        x = 0.7 + (index % 3) * 4.0
        y = 1.35 + (index // 3) * 1.35
        _add_placeholder_box(slide, x, y, 3.65, 1.05, item, fill=RGBColor(255, 255, 255))
    _add_placeholder_box(slide, 0.7, 4.65, 7.9, 0.85, "Implication / recommendation / board ask", fill=PALE_PEACH)


def _add_competitor_layout(slide, slide_plan: dict[str, Any]) -> None:
    add_table_placeholder(
        slide,
        0.7,
        1.25,
        7.95,
        3.95,
        ["Player / segment", "Proposition", "Strength", "Gap"],
        [
            ["Competitor A", "Placeholder", "Placeholder", "Placeholder"],
            ["Competitor B", "Placeholder", "Placeholder", "Placeholder"],
            ["Competitor C", "Placeholder", "Placeholder", "Placeholder"],
            ["OSK implication", "Placeholder", "Placeholder", "Placeholder"],
        ],
    )
    _add_placeholder_box(slide, 8.95, 1.25, 3.45, 2.0, "Key observations sidebar", fill=PALE_BLUE)
    _add_placeholder_box(slide, 8.95, 3.55, 3.45, 1.65, "Implication / recommendation", fill=PALE_PEACH)


def _add_market_landscape_layout(slide, slide_plan: dict[str, Any]) -> None:
    _add_placeholder_box(slide, 0.7, 1.25, 7.95, 4.0, "Segmented landscape / market map placeholder\n\nCategory 1 | Category 2 | Category 3", fill=RGBColor(255, 255, 255))
    _add_placeholder_box(slide, 8.95, 1.25, 3.45, 1.25, "Observation 1", fill=PALE_BLUE)
    _add_placeholder_box(slide, 8.95, 2.75, 3.45, 1.25, "Observation 2", fill=PALE_BLUE)
    _add_placeholder_box(slide, 8.95, 4.25, 3.45, 1.0, "So what for OSK", fill=PALE_PEACH)


def _add_market_sizing_layout(slide, slide_plan: dict[str, Any]) -> None:
    add_chart_placeholder(slide, 0.75, 1.25, 7.25, 3.9, "Chart placeholder: sizing / financial analysis")
    _add_placeholder_box(slide, 8.35, 1.25, 4.0, 1.55, "Assumptions box", fill=PALE_BLUE)
    _add_placeholder_box(slide, 8.35, 3.05, 4.0, 1.0, "Key takeaway", fill=PALE_PEACH)
    _add_placeholder_box(slide, 8.35, 4.3, 4.0, 0.85, "Data source / confidence notes", fill=LIGHT)


def _add_generic_layout(slide, slide_plan: dict[str, Any]) -> None:
    content = "\n".join(_content_or_defaults(slide_plan, ["Main content placeholder", "Supporting evidence placeholder"]))
    _add_placeholder_box(slide, 0.7, 1.25, 7.95, 3.95, content, fill=RGBColor(255, 255, 255))
    _add_placeholder_box(slide, 8.95, 1.25, 3.45, 1.75, "Supporting insight / callout", fill=PALE_BLUE)
    _add_placeholder_box(slide, 8.95, 3.35, 3.45, 1.85, "Implication / next step", fill=PALE_PEACH)


def _add_build_notes_box(slide, slide_plan: dict[str, Any]) -> None:
    notes = build_slide_notes(slide_plan)
    excerpt = notes[:420] + ("..." if len(notes) > 420 else "")
    _add_placeholder_box(slide, 0.7, 5.78, 11.75, 1.05, f"Build Notes\n{excerpt}", fill=RGBColor(250, 250, 250), font_size=7)


def _footer_text(slide_plan: dict[str, Any]) -> str:
    matched = slide_plan.get("matched_template") or {}
    if matched:
        return f"{matched.get('deck_name', 'Matched template')} slide {matched.get('slide_number', '?')} | {slide_plan.get('slide_type', 'other')}"
    return f"No matched template | {slide_plan.get('slide_type', 'other')}"


def _add_placeholder_box(
    slide,
    x: float,
    y: float,
    w: float,
    h: float,
    text: str,
    *,
    fill=LIGHT,
    border=BORDER,
    font_size: int = 10,
) -> Any:
    shape = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE,
        Inches(x),
        Inches(y),
        Inches(w),
        Inches(h),
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = border
    shape.line.width = Pt(0.7)
    _set_shape_text(shape, text, font_size=font_size, color=DARK)
    return shape


def _add_textbox(slide, x: float, y: float, w: float, h: float, text: str, font_size: int, *, bold: bool = False, color=DARK) -> Any:
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    _set_shape_text(shape, text, font_size=font_size, bold=bold, color=color)
    return shape


def _set_shape_text(shape, text: str, *, font_size: int, bold: bool = False, color=DARK) -> None:
    text_frame = shape.text_frame
    text_frame.clear()
    text_frame.word_wrap = True
    text_frame.margin_left = Inches(0.08)
    text_frame.margin_right = Inches(0.08)
    text_frame.margin_top = Inches(0.05)
    text_frame.margin_bottom = Inches(0.04)
    for index, line in enumerate(str(text).splitlines() or [""]):
        paragraph = text_frame.paragraphs[0] if index == 0 else text_frame.add_paragraph()
        paragraph.text = line
        paragraph.font.size = Pt(font_size if index else font_size)
        paragraph.font.name = "Aptos"
        paragraph.font.bold = bold if index == 0 else False
        paragraph.font.color.rgb = color
        paragraph.alignment = PP_ALIGN.LEFT


def _format_cell(cell, *, fill, color, bold: bool = False) -> None:
    cell.fill.solid()
    cell.fill.fore_color.rgb = fill
    text_frame = cell.text_frame
    for paragraph in text_frame.paragraphs:
        for run in paragraph.runs:
            run.font.name = "Aptos"
            run.font.size = Pt(8)
            run.font.bold = bold
            run.font.color.rgb = color


def _add_background(slide) -> None:
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = RGBColor(255, 255, 255)


def _content_or_defaults(slide_plan: dict[str, Any], defaults: list[str]) -> list[str]:
    content = slide_plan.get("suggested_content")
    if isinstance(content, list) and content:
        return [str(item) for item in content]
    if isinstance(content, str) and content.strip():
        return [content]
    return defaults
