from __future__ import annotations

import json
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from config import OUTPUT_DIRECTORIES
from pptx_generation.generation_models import GeneratedDeckInput, GeneratedDeckMetadata
from pptx_generation.notes_builder import build_slide_notes
from pptx_generation.pptx_repository import (
    default_output_path,
    ensure_generated_deck_dirs,
    get_latest_recommendation_path,
    save_generation_metadata,
    save_speaker_notes_markdown,
)
from recommendation.recommend_from_content import recommend_from_content


def load_recommendation_json(path: str | Path) -> dict[str, Any]:
    recommendation_path = Path(path)
    if not recommendation_path.exists() and recommendation_path.name == "latest_recommendation.json":
        latest = get_latest_recommendation_path()
        if latest:
            recommendation_path = latest
    if not recommendation_path.exists():
        raise FileNotFoundError(f"Recommendation JSON not found: {path}")
    try:
        payload = json.loads(recommendation_path.read_text(encoding="utf-8"))
    except json.JSONDecodeError as exc:
        raise ValueError(f"Invalid recommendation JSON: {recommendation_path}") from exc
    if not isinstance(payload, dict) or not payload:
        raise ValueError(f"Recommendation JSON is empty or malformed: {recommendation_path}")
    payload["_source_recommendation_path"] = str(recommendation_path)
    return payload


def normalize_recommendation_to_slide_plans(recommendation: dict[str, Any]) -> list[dict[str, Any]]:
    if not isinstance(recommendation, dict) or not recommendation:
        return []

    patterns = _extract_patterns(recommendation)
    matching_method = recommendation.get("matching_method") or _pattern_matching_method(patterns)
    plans: list[dict[str, Any]] = []
    used_pattern_indices: set[int] = set()

    deck_slides = _extract_deck_slides(recommendation)
    for index, item in enumerate(deck_slides, start=1):
        plan = _normalize_slide_item(
            item=item,
            index=index,
            recommendation=recommendation,
            patterns=patterns,
            used_pattern_indices=used_pattern_indices,
            matching_method=matching_method,
        )
        plans.append(plan)

    if plans:
        return plans

    single_slide = _extract_single_slide(recommendation)
    if single_slide:
        plans.append(
            _normalize_slide_item(
                item=single_slide,
                index=1,
                recommendation=recommendation,
                patterns=patterns,
                used_pattern_indices=used_pattern_indices,
                matching_method=matching_method,
            )
        )

    return plans


def generate_placeholder_deck_from_recommendation(
    recommendation_path: str | Path,
    output_path: str | Path | None = None,
) -> dict[str, Any]:
    recommendation = load_recommendation_json(recommendation_path)
    source_path = recommendation.get("_source_recommendation_path") or str(recommendation_path)
    return _generate_placeholder_deck(
        recommendation=recommendation,
        source_recommendation_path=source_path,
        output_path=output_path,
    )


def generate_placeholder_deck_from_text(
    text: str,
    goal: str | None,
    mode: str = "deck",
    output_path: str | Path | None = None,
) -> dict[str, Any]:
    recommendation = recommend_from_content(text=text, user_goal=goal, mode=mode, prompt_output=False)
    return _generate_placeholder_deck(
        recommendation=recommendation,
        source_recommendation_path=recommendation.get("recommendation_json_path"),
        output_path=output_path,
    )


def generate_placeholder_deck_from_file(
    file_path: str | Path,
    goal: str | None,
    mode: str = "deck",
    output_path: str | Path | None = None,
) -> dict[str, Any]:
    recommendation = recommend_from_content(
        file_path=str(file_path),
        user_goal=goal,
        mode=mode,
        prompt_output=False,
    )
    return _generate_placeholder_deck(
        recommendation=recommendation,
        source_recommendation_path=recommendation.get("recommendation_json_path"),
        output_path=output_path,
    )


def _generate_placeholder_deck(
    *,
    recommendation: dict[str, Any],
    source_recommendation_path: str | None,
    output_path: str | Path | None,
) -> dict[str, Any]:
    try:
        from pptx import Presentation
        from pptx.util import Inches
    except ImportError as exc:
        raise ImportError(
            "python-pptx is required for placeholder deck generation. "
            "Install it with: .venv/bin/python -m pip install python-pptx"
        ) from exc

    from pptx_generation.slide_layout_builder import (
        SLIDE_H,
        SLIDE_W,
        add_agenda_slide,
        add_appendix_slide,
        add_content_slide,
        add_cover_slide,
    )

    ensure_generated_deck_dirs()
    slide_plans = normalize_recommendation_to_slide_plans(recommendation)
    if not slide_plans:
        raise ValueError("No slide plans found in recommendation JSON.")

    timestamp = datetime.now(timezone.utc).strftime("%Y%m%d_%H%M%S")
    title = _deck_title(recommendation)
    target_path = Path(output_path) if output_path else default_output_path(title, timestamp)
    if not target_path.is_absolute():
        target_path = Path.cwd() / target_path
    target_path.parent.mkdir(parents=True, exist_ok=True)

    deck_input = GeneratedDeckInput(
        title=title,
        user_goal=_user_goal(recommendation),
        mode=recommendation.get("recommended_output_mode", "deck"),
        source_recommendation_path=source_recommendation_path,
        matching_method=recommendation.get("matching_method") or _pattern_matching_method(_extract_patterns(recommendation)),
        slide_patterns=_extract_patterns(recommendation),
        recommendation=recommendation,
    ).as_dict()

    prs = Presentation()
    prs.slide_width = Inches(SLIDE_W)
    prs.slide_height = Inches(SLIDE_H)
    add_cover_slide(prs, deck_input)
    add_agenda_slide(prs, slide_plans)
    for index, slide_plan in enumerate(slide_plans, start=1):
        add_content_slide(prs, slide_plan, index)
    add_appendix_slide(prs, deck_input, slide_plans)
    prs.save(target_path)

    notes_path = save_speaker_notes_markdown(
        [
            {
                "slide_number": str(index),
                "slide_title": slide_plan.get("slide_title", f"Slide {index}"),
                "notes": build_slide_notes(slide_plan),
            }
            for index, slide_plan in enumerate(slide_plans, start=1)
        ],
        target_path,
    )

    metadata = GeneratedDeckMetadata(
        generated_at=datetime.now(timezone.utc).isoformat(),
        output_path=str(target_path),
        source_recommendation_path=source_recommendation_path,
        user_goal=deck_input.get("user_goal", ""),
        mode=deck_input.get("mode", "deck"),
        number_of_slides=len(prs.slides),
        number_of_content_slides=len(slide_plans),
        matching_method=deck_input.get("matching_method", "metadata_match"),
        referenced_template_slides=_referenced_template_slides(slide_plans),
        warnings=_warnings(recommendation, slide_plans),
    ).as_dict()
    metadata["speaker_notes_path"] = str(notes_path)
    metadata_path = save_generation_metadata(metadata, target_path)
    metadata["metadata_path"] = str(metadata_path)

    log_path = OUTPUT_DIRECTORIES["generated_deck_logs"] / f"{target_path.stem}.log"
    log_path.write_text(
        f"Generated {target_path}\nSlides: {len(prs.slides)}\nMetadata: {metadata_path}\n",
        encoding="utf-8",
    )
    metadata["log_path"] = str(log_path)
    return metadata


def _extract_deck_slides(recommendation: dict[str, Any]) -> list[dict[str, Any]]:
    candidates = [
        recommendation.get("recommended_slide_sequence"),
        recommendation.get("slides"),
        recommendation.get("slide_recommendations"),
        recommendation.get("recommendation", {}).get("slides")
        if isinstance(recommendation.get("recommendation"), dict)
        else None,
        recommendation.get("recommendation", {}).get("recommended_slide_sequence")
        if isinstance(recommendation.get("recommendation"), dict)
        else None,
    ]
    for candidate in candidates:
        normalized = _list_of_dicts(candidate)
        if normalized:
            return normalized

    for key in ["recommended_deck_structure", "deck_structure", "recommended_structure"]:
        structure = recommendation.get(key)
        if not isinstance(structure, dict):
            continue
        slides: list[dict[str, Any]] = []
        for section in _list_of_dicts(structure.get("sections")):
            section_name = section.get("section_name") or ""
            section_purpose = section.get("section_purpose") or ""
            for slide in _list_of_dicts(section.get("recommended_slides") or section.get("slides")):
                item = dict(slide)
                item["_section_name"] = section_name
                item["_section_purpose"] = section_purpose
                slides.append(item)
        if slides:
            return slides

    return []


def _extract_single_slide(recommendation: dict[str, Any]) -> dict[str, Any] | None:
    for key in ["recommended_slide_structure", "slide_structure", "recommended_slide"]:
        value = recommendation.get(key)
        if isinstance(value, dict) and value:
            item = dict(value)
            item.setdefault("slide_title", recommendation.get("recommended_slide_or_deck_title"))
            item.setdefault("slide_type", item.get("recommended_slide_type"))
            item.setdefault("chart_type", item.get("recommended_chart_type"))
            item.setdefault("storyline_type", item.get("recommended_storyline_type"))
            return item
    return None


def _normalize_slide_item(
    *,
    item: dict[str, Any],
    index: int,
    recommendation: dict[str, Any],
    patterns: list[dict[str, Any]],
    used_pattern_indices: set[int],
    matching_method: str,
) -> dict[str, Any]:
    slide_type = _first_text(
        item.get("slide_type"),
        item.get("recommended_slide_type"),
        recommendation.get("recommended_slide_structure", {}).get("recommended_slide_type")
        if isinstance(recommendation.get("recommended_slide_structure"), dict)
        else None,
        "other",
    )
    chart_type = _first_text(
        item.get("chart_type"),
        item.get("recommended_chart_type"),
        recommendation.get("recommended_slide_structure", {}).get("recommended_chart_type")
        if isinstance(recommendation.get("recommended_slide_structure"), dict)
        else None,
        "unknown",
    )
    title = _first_text(
        item.get("suggested_action_title"),
        item.get("slide_title"),
        recommendation.get("recommended_slide_or_deck_title"),
        f"Placeholder slide {index}",
    )
    content_blocks = _string_list(
        item.get("content_blocks")
        or item.get("suggested_content")
        or item.get("content_to_use_from_source")
        or item.get("template_instruction")
    )
    matched_template = _extract_matched_template(item) or _best_matching_pattern(
        patterns,
        slide_type,
        index,
        used_pattern_indices,
    )
    return {
        "slide_number": item.get("slide_number") or index,
        "slide_title": title,
        "slide_type": slide_type,
        "storyline": _first_text(
            item.get("storyline"),
            item.get("storyline_type"),
            recommendation.get("recommended_storyline"),
            "",
        ),
        "slide_objective": _first_text(
            item.get("slide_objective"),
            item.get("_section_purpose"),
            item.get("content_to_use_from_source"),
            "Create an editable placeholder slide based on the recommendation.",
        ),
        "layout_pattern": _first_text(
            item.get("layout_pattern"),
            item.get("layout_instruction"),
            item.get("template_instruction"),
            recommendation.get("recommended_slide_structure", {}).get("layout_instruction")
            if isinstance(recommendation.get("recommended_slide_structure"), dict)
            else None,
            "",
        ),
        "chart_type": chart_type,
        "business_context": _first_text(
            item.get("business_context"),
            recommendation.get("content_diagnosis"),
            "",
        ),
        "suggested_content": content_blocks,
        "tags": _string_list(item.get("tags")),
        "matched_template": matched_template,
        "matching_method": _first_text(item.get("matching_method"), matching_method, "metadata_match"),
        "data_required": _string_list(item.get("data_required")),
        "analysis_required": _string_list(item.get("analysis_required")),
        "missing_data_or_research_gaps": _string_list(
            item.get("missing_data_or_research_gaps") or item.get("additional_data_needed")
        ),
    }


def _extract_patterns(recommendation: dict[str, Any]) -> list[dict[str, Any]]:
    for key in [
        "similar_slide_patterns_found",
        "matched_slide_patterns",
        "similar_slide_patterns",
        "relevant_slide_patterns",
        "top_matched_slide_patterns",
    ]:
        patterns = _list_of_dicts(recommendation.get(key))
        if patterns:
            return patterns
    return []


def _best_matching_pattern(
    patterns: list[dict[str, Any]],
    slide_type: str,
    index: int,
    used_pattern_indices: set[int],
) -> dict[str, Any]:
    if not patterns:
        return {}
    for pattern_index, pattern in enumerate(patterns):
        if pattern_index in used_pattern_indices:
            continue
        if pattern.get("slide_type") == slide_type:
            used_pattern_indices.add(pattern_index)
            return dict(pattern)
    fallback_index = min(index - 1, len(patterns) - 1)
    used_pattern_indices.add(fallback_index)
    return dict(patterns[fallback_index])


def _extract_matched_template(item: dict[str, Any]) -> dict[str, Any]:
    for key in ["matched_template", "reference_slide_metadata", "matched_slide_pattern"]:
        value = item.get(key)
        if isinstance(value, dict):
            return dict(value)
    return {}


def _deck_title(recommendation: dict[str, Any]) -> str:
    deck_structure = recommendation.get("recommended_deck_structure")
    if isinstance(deck_structure, dict):
        title = deck_structure.get("suggested_deck_title")
        if title:
            return str(title)
    return _first_text(
        recommendation.get("recommended_slide_or_deck_title"),
        recommendation.get("user_goal"),
        "Generated Placeholder Deck",
    )


def _user_goal(recommendation: dict[str, Any]) -> str:
    return _first_text(
        recommendation.get("user_goal"),
        recommendation.get("goal"),
        recommendation.get("brief_interpretation"),
        "",
    )


def _referenced_template_slides(slide_plans: list[dict[str, Any]]) -> list[dict[str, Any]]:
    references = []
    seen = set()
    for plan in slide_plans:
        matched = plan.get("matched_template") or {}
        if not matched:
            continue
        key = (matched.get("deck_name"), matched.get("slide_number"))
        if key in seen:
            continue
        seen.add(key)
        references.append(
            {
                "deck_name": matched.get("deck_name"),
                "slide_number": matched.get("slide_number"),
                "slide_title": matched.get("slide_title"),
                "slide_type": matched.get("slide_type"),
                "chart_type": matched.get("chart_type"),
                "similarity_score": matched.get("similarity_score"),
                "match_score": matched.get("match_score"),
            }
        )
    return references


def _warnings(recommendation: dict[str, Any], slide_plans: list[dict[str, Any]]) -> list[str]:
    warnings = ["Placeholder deck only; final design automation and OSK style replication are not implemented in Stage 7."]
    if not _extract_patterns(recommendation):
        warnings.append("No matched slide patterns were found in the recommendation output.")
    if any(not (plan.get("matched_template") or {}) for plan in slide_plans):
        warnings.append("One or more content slides do not have a matched template reference.")
    return warnings


def _pattern_matching_method(patterns: list[dict[str, Any]]) -> str:
    if not patterns:
        return "metadata_match"
    return str(patterns[0].get("matching_method") or "metadata_match")


def _list_of_dicts(value: object) -> list[dict[str, Any]]:
    if not isinstance(value, list):
        return []
    return [dict(item) for item in value if isinstance(item, dict)]


def _first_text(*values: object) -> str:
    for value in values:
        if value is None:
            continue
        text = str(value).strip()
        if text:
            return text
    return ""


def _string_list(value: object) -> list[str]:
    if value is None:
        return []
    if isinstance(value, str):
        return [value] if value.strip() else []
    if isinstance(value, list):
        return [str(item).strip() for item in value if str(item).strip()]
    return [str(value).strip()] if str(value).strip() else []
