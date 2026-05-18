from __future__ import annotations

import json
from collections import Counter
from pathlib import Path
from typing import Any

from quality_control.deck_qa_models import (
    ScoreBreakdown,
    add_recommendation_once,
    clamp_category_score,
    find_duplicate_titles,
    find_near_duplicate_titles,
    issue,
    now_iso,
    rating_for_score,
    repeated_template_usage,
    safe_path,
    word_count,
)


def review_generated_deck(
    pptx_path: str | Path,
    metadata_path: str | Path | None = None,
    notes_path: str | Path | None = None,
) -> dict[str, Any]:
    path = safe_path(pptx_path)
    issues: list[dict[str, Any]] = []
    recommendations: list[str] = []
    slide_level_findings: list[dict[str, Any]] = []
    metadata = _load_optional_json(metadata_path)
    notes = _load_optional_text(notes_path)

    if path is None or not path.exists():
        issues.append(
            issue(
                "critical",
                "generation",
                f"PPTX file is missing or unreadable: {pptx_path}",
                "Regenerate the deck or provide the correct PPTX path.",
            )
        )
        score = 0
        return _result(path, metadata_path, notes_path, score, {}, issues, recommendations, slide_level_findings)

    try:
        from pptx import Presentation

        prs = Presentation(str(path))
    except Exception as exc:
        issues.append(
            issue(
                "critical",
                "generation",
                f"PPTX file could not be opened: {exc}",
                "Regenerate the deck or inspect the file manually in PowerPoint.",
            )
        )
        score = 0
        return _result(path, metadata_path, notes_path, score, {}, issues, recommendations, slide_level_findings)

    slide_infos = [_inspect_slide(slide, index) for index, slide in enumerate(prs.slides, start=1)]
    total_slides = len(slide_infos)
    cover_found = any(info["role"] == "cover" for info in slide_infos)
    agenda_found = any(info["role"] == "agenda" for info in slide_infos)
    appendix_found = any(info["role"] == "appendix" for info in slide_infos)
    content_slide_count = _metadata_int(metadata, "number_of_content_slides")
    if content_slide_count is None:
        content_slide_count = sum(1 for info in slide_infos if info["role"] == "content")

    _check_structure(
        total_slides,
        content_slide_count,
        cover_found,
        agenda_found,
        appendix_found,
        issues,
        recommendations,
    )
    _check_metadata_and_notes(metadata, notes, metadata_path, notes_path, issues, recommendations)
    _check_generation_metadata(metadata, content_slide_count, issues, recommendations)
    _check_slide_texts(slide_infos, issues, recommendations, slide_level_findings)
    _check_template_metadata(metadata, content_slide_count, issues, recommendations)

    score_breakdown = _score_generated_deck(
        total_slides=total_slides,
        content_slide_count=content_slide_count,
        cover_found=cover_found,
        agenda_found=agenda_found,
        appendix_found=appendix_found,
        metadata=metadata,
        notes=notes,
        issues=issues,
    )
    overall_score = score_breakdown.total
    summary = {
        "total_slides": total_slides,
        "content_slides": content_slide_count,
        "cover_found": cover_found,
        "agenda_found": agenda_found,
        "appendix_found": appendix_found,
        "metadata_found": bool(metadata),
        "notes_found": bool(notes),
        "cloned_slide_count": _metadata_int(metadata, "content_slides_cloned_from_source_pptx"),
        "image_fallback_count": _metadata_int(metadata, "content_slides_with_image_fallback"),
        "matching_method": metadata.get("matching_method") if metadata else None,
        "referenced_template_count": len(metadata.get("referenced_template_slides", [])) if metadata else 0,
        "blank_or_near_empty_slide_count": sum(1 for info in slide_infos if info["is_near_empty"]),
        "dense_slide_count": sum(1 for info in slide_infos if info["is_too_dense"]),
    }
    return {
        "qa_type": "generated_deck",
        "reviewed_path": str(path),
        "metadata_path": str(safe_path(metadata_path)) if metadata_path else None,
        "notes_path": str(safe_path(notes_path)) if notes_path else None,
        "generated_at": now_iso(),
        "overall_score": overall_score,
        "rating": rating_for_score(overall_score),
        "score_breakdown": score_breakdown.as_dict(),
        "summary": summary,
        "issues": issues,
        "recommendations": recommendations,
        "slide_level_findings": slide_level_findings,
        "metadata": metadata,
    }


def _inspect_slide(slide, index: int) -> dict[str, Any]:
    text_runs = []
    for shape in slide.shapes:
        text = getattr(shape, "text", "") if hasattr(shape, "text") else ""
        if text and text.strip():
            text_runs.append(text.strip())
    all_text = "\n".join(text_runs)
    lower_text = all_text.lower()
    role = "content"
    if index == 1 or "generated placeholder deck" in lower_text or "reference-style draft" in lower_text:
        role = "cover"
    elif "agenda" in lower_text[:500]:
        role = "agenda"
    elif "appendix" in lower_text[:500] or "matched template patterns" in lower_text:
        role = "appendix"
    title = _best_title(text_runs)
    char_count = len(all_text)
    return {
        "slide_number": index,
        "role": role,
        "title": title,
        "text_character_count": char_count,
        "shape_count": len(slide.shapes),
        "text_shape_count": len(text_runs),
        "is_near_empty": role == "content" and char_count < 30,
        "is_too_dense": role == "content" and char_count > 1200,
    }


def _best_title(text_runs: list[str]) -> str:
    candidates = []
    for text in text_runs[:8]:
        first_line = text.splitlines()[0].strip()
        if not first_line:
            continue
        if " · " in first_line and len(first_line) <= 28:
            continue
        if first_line.lower().startswith(("editable content", "editable build note", "reference background")):
            continue
        candidates.append(first_line)
    if not candidates:
        return text_runs[0].splitlines()[0].strip() if text_runs else ""
    # The title is normally the first prominent text shape. Returning the
    # longest early text run mistakes subtitles, callouts, or agenda rows for
    # slide titles, especially in content-first generated decks.
    return candidates[0]


def _check_structure(
    total_slides: int,
    content_slide_count: int,
    cover_found: bool,
    agenda_found: bool,
    appendix_found: bool,
    issues: list[dict[str, Any]],
    recommendations: list[str],
) -> None:
    if total_slides == 0:
        issues.append(
            issue("critical", "structure", "Generated deck is empty.", "Regenerate the deck.")
        )
    if content_slide_count == 0:
        issues.append(
            issue("critical", "structure", "No content slides detected.", "Regenerate from a valid recommendation.")
        )
    if not cover_found:
        issues.append(issue("medium", "structure", "Cover slide was not detected.", "Add or repair a cover slide."))
    if not agenda_found:
        issues.append(issue("medium", "structure", "Agenda slide was not detected.", "Add or repair the agenda slide."))
    if not appendix_found:
        issues.append(issue("low", "structure", "Appendix/reference slide was not detected.", "Add an appendix with template references and limitations."))
    if content_slide_count > 25:
        issues.append(
            issue(
                "high",
                "structure",
                f"Deck has {content_slide_count} content slides.",
                "Condense the core deck and move supporting pages to appendix.",
            )
        )
    elif content_slide_count > 20:
        issues.append(
            issue(
                "medium",
                "structure",
                f"Deck has {content_slide_count} content slides, which may be too long for a board-ready first draft.",
                "Condense to 10–15 core slides plus appendix.",
            )
        )
        add_recommendation_once(
            recommendations,
            "Condense the main storyline to 10–15 core slides and move detailed support to appendix.",
        )


def _check_metadata_and_notes(
    metadata: dict[str, Any],
    notes: str,
    metadata_path: str | Path | None,
    notes_path: str | Path | None,
    issues: list[dict[str, Any]],
    recommendations: list[str],
) -> None:
    if not metadata:
        issues.append(
            issue(
                "high",
                "generation",
                "No generation metadata file found or metadata could not be parsed.",
                "Provide the matching metadata JSON so QA can inspect clone/fallback counts and template references.",
            )
        )
    if not notes:
        issues.append(
            issue(
                "high",
                "generation",
                "No speaker notes/build notes markdown found or notes file is empty.",
                "Generate or attach speaker notes so consultants can review build instructions.",
            )
        )
    else:
        issues.append(
            issue(
                "info",
                "generation",
                "Speaker notes/build notes found.",
                "Use the notes file during manual consultant review.",
            )
        )


def _check_generation_metadata(
    metadata: dict[str, Any],
    content_slide_count: int,
    issues: list[dict[str, Any]],
    recommendations: list[str],
) -> None:
    if not metadata:
        return
    cloned = _metadata_int(metadata, "content_slides_cloned_from_source_pptx") or 0
    fallback = _metadata_int(metadata, "content_slides_with_image_fallback") or 0
    if content_slide_count and fallback / content_slide_count > 0.30:
        issues.append(
            issue(
                "high",
                "generation",
                f"{fallback}/{content_slide_count} content slides used image fallback.",
                "Review template source availability or improve clone handling.",
            )
        )
    elif fallback > 0:
        issues.append(
            issue(
                "medium",
                "generation",
                f"{fallback} content slide(s) used image fallback.",
                "Manually review fallback slides because their background is not editable.",
            )
        )
    if content_slide_count and cloned == content_slide_count and fallback == 0:
        issues.append(
            issue(
                "info",
                "generation",
                "All content slides cloned successfully from source PPTX.",
                "Still manually QA complex charts, groups, and master-layout behavior.",
            )
        )
    if metadata.get("warnings"):
        issues.append(
            issue(
                "low",
                "generation",
                f"Generator metadata contains {len(metadata.get('warnings', []))} warning(s).",
                "Read metadata warnings before using the deck.",
            )
        )
        add_recommendation_once(
            recommendations,
            "Review metadata warnings before sharing the deck.",
        )


def _check_slide_texts(
    slide_infos: list[dict[str, Any]],
    issues: list[dict[str, Any]],
    recommendations: list[str],
    slide_level_findings: list[dict[str, Any]],
) -> None:
    content_infos = [info for info in slide_infos if info["role"] == "content"]
    titles = [info["title"] for info in content_infos]
    for info in content_infos:
        count = word_count(info["title"])
        finding = {
            "slide_number": info["slide_number"],
            "title": info["title"],
            "word_count": count,
            "text_character_count": info["text_character_count"],
            "role": info["role"],
        }
        if info["is_near_empty"]:
            issues.append(
                issue(
                    "high",
                    "slide_quality",
                    "Content slide appears near-empty.",
                    "Populate or remove the slide.",
                    slide_number=info["slide_number"],
                )
            )
            finding["issue"] = "near_empty"
        if info["is_too_dense"]:
            issues.append(
                issue(
                    "medium",
                    "slide_quality",
                    "Content slide is text-dense.",
                    "Split content, simplify labels, or move detail to appendix.",
                    slide_number=info["slide_number"],
                )
            )
            finding["issue"] = "too_dense"
        if count > 25:
            issues.append(
                issue(
                    "high",
                    "slide_quality",
                    f"Slide title is very long ({count} words).",
                    "Shorten the title while preserving the action message.",
                    slide_number=info["slide_number"],
                )
            )
            finding["issue"] = "very_long_title"
        elif count > 18:
            issues.append(
                issue(
                    "medium",
                    "slide_quality",
                    f"Slide title is long ({count} words).",
                    "Tighten the action title to under 18 words.",
                    slide_number=info["slide_number"],
                )
            )
            finding["issue"] = "long_title"
        slide_level_findings.append(finding)

    duplicates = find_duplicate_titles(titles)
    if duplicates:
        issues.append(
            issue(
                "medium",
                "slide_quality",
                f"{len(duplicates)} duplicate title occurrence(s) detected in content slides.",
                "Merge duplicate slides or rewrite titles to clarify each slide's role.",
            )
        )
    near_duplicates = find_near_duplicate_titles(titles)
    if near_duplicates:
        issues.append(
            issue(
                "medium",
                "slide_quality",
                f"{len(near_duplicates)} near-duplicate content slide title pair(s) detected.",
                "Review the flow for repetitive arguments.",
            )
        )
        add_recommendation_once(recommendations, "Review duplicate or near-duplicate titles for repetitive deck flow.")


def _check_template_metadata(
    metadata: dict[str, Any],
    content_slide_count: int,
    issues: list[dict[str, Any]],
    recommendations: list[str],
) -> None:
    if not metadata:
        return
    matching_method = metadata.get("matching_method")
    if matching_method in {"vector_search", "per_slide_vector_search"}:
        issues.append(
            issue(
                "info",
                "template_matching",
                f"{matching_method} used for template matching.",
                "Keep embedding index available for better template retrieval.",
            )
        )
    else:
        issues.append(
            issue(
                "low",
                "template_matching",
                f"Matching method is {matching_method or 'unknown'}.",
                "Confirm semantic vector search was available when the recommendation was produced.",
            )
        )

    references = metadata.get("cloned_reference_slides") or metadata.get("referenced_template_slides") or []
    if not references:
        issues.append(
            issue(
                "high",
                "template_matching",
                "No referenced template slides found in metadata.",
                "Regenerate with template matching metadata enabled.",
            )
        )
        return

    for (deck_name, slide_number), count in repeated_template_usage(references):
        if count > 6:
            severity = "high"
        elif count > 3:
            severity = "medium"
        else:
            continue
        issues.append(
            issue(
                severity,
                "template_matching",
                f"Template slide reused {count} times: {deck_name} slide {slide_number}.",
                "Diversify matched reference slides or consolidate repeated content into appendix.",
            )
        )
        add_recommendation_once(recommendations, "Diversify overused source templates before polishing the deck.")

    scored = [ref for ref in metadata.get("referenced_template_slides", []) if ref.get("similarity_score") is not None]
    for ref in scored:
        try:
            score = float(ref.get("similarity_score"))
        except (TypeError, ValueError):
            continue
        if score < 0.35:
            severity = "high"
        elif score < 0.45:
            severity = "medium"
        else:
            continue
        issues.append(
            issue(
                severity,
                "template_matching",
                f"Low template similarity score ({score:.2f}) for {ref.get('deck_name')} slide {ref.get('slide_number')}.",
                "Manually review whether this reference slide is truly relevant.",
            )
        )


def _score_generated_deck(
    *,
    total_slides: int,
    content_slide_count: int,
    cover_found: bool,
    agenda_found: bool,
    appendix_found: bool,
    metadata: dict[str, Any],
    notes: str,
    issues: list[dict[str, Any]],
) -> ScoreBreakdown:
    structure = 25
    if not cover_found:
        structure -= 6
    if not agenda_found:
        structure -= 6
    if not appendix_found:
        structure -= 3
    if content_slide_count > 25:
        structure -= 9
    elif content_slide_count > 20:
        structure -= 5
    if total_slides == 0 or content_slide_count == 0:
        structure = 0

    # Cloned reference decks often carry source-slide text plus editable overlay
    # text, so density/title issues should prompt review without swamping the
    # whole score when structure and generation succeeded.
    slide_issue_count = sum(1 for item in issues if item.get("category") == "slide_quality")
    slide_quality = 25 - min(12, slide_issue_count)

    template_matching = 25
    if metadata.get("matching_method") not in {"vector_search", "per_slide_vector_search"}:
        template_matching -= 4
    if not (metadata.get("referenced_template_slides") or metadata.get("cloned_reference_slides")):
        template_matching -= 12
    template_matching -= sum(3 for item in issues if item.get("category") == "template_matching" and item.get("severity") in {"medium", "high"})

    generation = 25
    if not metadata:
        generation -= 12
    if not notes:
        generation -= 8
    fallback = _metadata_int(metadata, "content_slides_with_image_fallback") or 0
    if content_slide_count:
        generation -= min(10, round((fallback / content_slide_count) * 20))
    if metadata.get("clone_failures"):
        generation -= min(8, len(metadata.get("clone_failures", [])) * 2)

    return ScoreBreakdown(
        structure_quality=clamp_category_score(structure),
        slide_quality=clamp_category_score(slide_quality),
        template_matching_quality=clamp_category_score(template_matching),
        generation_quality=clamp_category_score(generation),
    )


def _result(
    path: Path | None,
    metadata_path: str | Path | None,
    notes_path: str | Path | None,
    score: int,
    summary: dict[str, Any],
    issues: list[dict[str, Any]],
    recommendations: list[str],
    slide_level_findings: list[dict[str, Any]],
) -> dict[str, Any]:
    return {
        "qa_type": "generated_deck",
        "reviewed_path": str(path) if path else None,
        "metadata_path": str(safe_path(metadata_path)) if metadata_path else None,
        "notes_path": str(safe_path(notes_path)) if notes_path else None,
        "generated_at": now_iso(),
        "overall_score": score,
        "rating": rating_for_score(score),
        "score_breakdown": ScoreBreakdown(0, 0, 0, 0).as_dict(),
        "summary": summary,
        "issues": issues,
        "recommendations": recommendations,
        "slide_level_findings": slide_level_findings,
    }


def _load_optional_json(path: str | Path | None) -> dict[str, Any]:
    if not path:
        return {}
    resolved = safe_path(path)
    if not resolved or not resolved.exists():
        return {}
    try:
        payload = json.loads(resolved.read_text(encoding="utf-8"))
    except json.JSONDecodeError:
        return {}
    return payload if isinstance(payload, dict) else {}


def _load_optional_text(path: str | Path | None) -> str:
    if not path:
        return ""
    resolved = safe_path(path)
    if not resolved or not resolved.exists():
        return ""
    return resolved.read_text(encoding="utf-8")


def _metadata_int(metadata: dict[str, Any], key: str) -> int | None:
    if not metadata or metadata.get(key) is None:
        return None
    try:
        return int(metadata.get(key))
    except (TypeError, ValueError):
        return None
