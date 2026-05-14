import csv
import sys
import tempfile
import unittest
from pathlib import Path


PROJECT_ROOT = Path(__file__).resolve().parents[1]
SRC_DIR = PROJECT_ROOT / "src"
if str(SRC_DIR) not in sys.path:
    sys.path.insert(0, str(SRC_DIR))

from content.extract_content import extract_content


class ContentExtractionTests(unittest.TestCase):
    def test_txt_extraction_works(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "notes.txt"
            path.write_text("Market entry notes\nCompetitor A is active.", encoding="utf-8")

            result = extract_content(file_path=str(path))

            self.assertEqual(result["source_type"], "txt")
            self.assertIn("Market entry notes", result["raw_text"])
            self.assertTrue(Path(result["raw_text_path"]).exists())

    def test_pasted_text_works(self):
        result = extract_content(text="Malaysia EWA market notes")

        self.assertEqual(result["source_type"], "text")
        self.assertEqual(result["source_name"], "pasted_text")
        self.assertIn("EWA", result["raw_text"])

    def test_docx_extraction_works(self):
        from docx import Document

        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "sample.docx"
            document = Document()
            document.add_paragraph("Board strategy update")
            table = document.add_table(rows=2, cols=2)
            table.cell(0, 0).text = "Metric"
            table.cell(0, 1).text = "Value"
            table.cell(1, 0).text = "Revenue"
            table.cell(1, 1).text = "100"
            document.save(path)

            result = extract_content(file_path=str(path))

            self.assertEqual(result["source_type"], "docx")
            self.assertIn("Board strategy update", result["raw_text"])
            self.assertIn("Revenue | 100", result["raw_text"])

    def test_csv_extraction_works(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "sample.csv"
            with path.open("w", newline="", encoding="utf-8") as handle:
                writer = csv.writer(handle)
                writer.writerow(["company", "revenue"])
                writer.writerow(["A", "10"])
                writer.writerow(["B", "12"])

            result = extract_content(file_path=str(path))

            self.assertEqual(result["source_type"], "csv")
            self.assertIn("Column names: company, revenue", result["raw_text"])


if __name__ == "__main__":
    unittest.main()
