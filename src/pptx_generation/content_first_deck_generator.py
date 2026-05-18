from __future__ import annotations

import json
import os
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from config import OUTPUT_DIRECTORIES
from content.analyze_content import analyze_content
from pptx_generation.content_first_layout_builder import (
    SLIDE_H,
    SLIDE_W,
    add_consultant_agenda,
    add_consultant_cover,
    add_planned_content_slide,
    add_source_grounding_appendix,
)
from pptx_generation.notes_builder import build_slide_notes
from pptx_generation.pptx_repository import (
    default_output_path,
    ensure_generated_deck_dirs,
    save_generation_metadata,
    save_speaker_notes_markdown,
)
from pptx_generation.source_slide_copier import SlideCloneError, clone_reference_slide_as_content_shell
from recommendation.match_slide_patterns import find_relevant_patterns_for_planned_slide
from slide_planning.generate_slide_plan import generate_slide_plan
from utils.json_utils import write_json


DEFAULT_TEMPLATE_FIT_THRESHOLD = 0.62


def generate_content_first_deck_from_text(
    text: str,
    goal: str | None,
    mode: str = "deck",
    output_path: str | Path | None = None,
) -> dict[str, Any]:
    content_analysis = analyze_content(text=text, user_goal=goal, mode=mode)
    plan = generate_slide_plan(content_analysis=content_analysis, user_goal=goal, mode=mode)
    return generate_content_first_deck_from_plan(
        slide_plan=plan,
        content_analysis=content_analysis,
        output_path=output_path,
    )


def generate_content_first_deck_from_file(
    file_path: str | Path,
    goal: str | None,
    mode: str = "deck",
    output_path: str | Path | None = None,
) -> dict[str, Any]:
    content_analysis = analyze_content(file_path=str(file_path), user_goal=goal, mode=mode)
    plan = generate_slide_plan(content_analysis=content_analysis, user_goal=goal, mode=mode)
    return generate_content_first_deck_from_plan(
        slide_plan=plan,
        content_analysis=content_analysis,
        output_path=output_path,
    )


def generate_content_first_deck_from_plan(
    *,
    slide_plan: dict[str, Any],
    content_analysis: dict[str, Any] | None = None,
    output_path: str | Path | None = None,
) -> dict[str, Any]:
    try:
        from pptx import Presentation
        from pptx.util import Inches
    except ImportError as exc:
        raise ImportError("python-pptx is required for content-first deck generation.") from exc

    ensure_generated_deck_dirs()
    slides = _prepare_planned_slides(slide_plan, content_analysis or {})
    if not slides:
        raise ValueError("Content-first slide plan does not contain any slides.")

    timestamp = datetime.now(timezone.utc).strftime("%Y%m%d_%H%M%S")
    title = str(slide_plan.get("deck_title") or "Content-first board-ready deck")
    target_path = Path(output_path) if output_path else default_output_path(title, timestamp)
    if not target_path.is_absolute():
        target_path = Path.cwd() / target_path
    target_path.parent.mkdir(parents=True, exist_ok=True)

    prs = Presentation()
    prs.slide_width = Inches(SLIDE_W)
    prs.slide_height = Inches(SLIDE_H)
    add_consultant_cover(prs, slide_plan)
    add_consultant_agenda(prs, slides)

    presentation_cache: dict[str, Any] = {}
    template_shell_count = 0
    custom_layout_count = 0
    clone_failures: list[dict[str, Any]] = []
    referenced_templates: list[dict[str, Any]] = []

    for index, planned_slide in enumerate(slides, start=1):
        if planned_slide.get("generation_strategy") == "template_shell":
            try:
                _, reference = clone_reference_slide_as_content_shell(prs, planned_slide, presentation_cache)
                template_shell_count += 1
                referenced_templates.append(_reference_item(planned_slide, reference))
                continue
            except Exception as exc:
                clone_failures.append(
                    {
                        "slide_number": planned_slide.get("slide_number"),
                        "slide_title": planned_slide.get("slide_title"),
                        "error": str(exc),
                        "matched_template": planned_slide.get("matched_template"),
                    }
                )
                planned_slide["generation_strategy"] = "custom_layout_after_template_clone_failure"
        add_planned_content_slide(prs, planned_slide, index)
        custom_layout_count += 1

    slide_plan["generation_mode"] = "content_first_consultant"
    slide_plan["template_shell_count"] = template_shell_count
    slide_plan["custom_layout_count"] = custom_layout_count
    add_source_grounding_appendix(prs, slide_plan, slides)
    prs.save(target_path)

    notes_path = save_speaker_notes_markdown(_speaker_notes(slides), target_path)
    grounding_path = _save_source_grounding_markdown(slide_plan, slides, target_path)
    enriched_plan_path = _save_enriched_slide_plan(slide_plan, slides, target_path)

    metadata = {
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "output_path": str(target_path),
        "user_goal": slide_plan.get("user_goal") or "",
        "mode": slide_plan.get("output_mode") or slide_plan.get("requested_mode") or "deck",
        "number_of_slides": len(prs.slides),
        "number_of_content_slides": len(slides),
        "matching_method": _matching_method(slides),
        "generator_version": "stage9_content_first_consultant_v1",
        "generation_mode": "content_first_consultant",
        "content_sufficiency_score": slide_plan.get("content_sufficiency_score"),
        "content_sufficiency_level": slide_plan.get("content_sufficiency_level"),
        "slide_plan_json_path": enriched_plan_path,
        "speaker_notes_path": str(notes_path),
        "source_grounding_report_path": str(grounding_path),
        "referenced_template_slides": _dedupe_references(referenced_templates + _candidate_reference_items(slides)),
        "template_shell_count": template_shell_count,
        "custom_layout_count": custom_layout_count,
        "content_slides_cloned_from_source_pptx": template_shell_count,
        "content_slides_with_image_fallback": 0,
        "clone_failures": clone_failures,
        "slide_outline": [
            {
                "slide_number": slide.get("slide_number"),
                "slide_title": slide.get("slide_title"),
                "slide_type": slide.get("slide_type"),
                "generation_strategy": slide.get("generation_strategy"),
                "matched_template": slide.get("matched_template"),
            }
            for slide in slides
        ],
        "warnings": list(slide_plan.get("generation_warnings") or [])
        + [
            "Content-first generator creates source-grounded editable slides. Template shells are used only when the per-slide fit is strong enough.",
        ],
    }
    metadata_path = save_generation_metadata(metadata, target_path)
    metadata["metadata_path"] = str(metadata_path)
    log_path = OUTPUT_DIRECTORIES["generated_deck_logs"] / f"{target_path.stem}.log"
    log_path.write_text(
        "\n".join(
            [
                f"Generated content-first deck: {target_path}",
                f"Slides: {len(prs.slides)}",
                f"Content slides: {len(slides)}",
                f"Template shells: {template_shell_count}",
                f"Custom layouts: {custom_layout_count}",
                f"Metadata: {metadata_path}",
            ]
        )
        + "\n",
        encoding="utf-8",
    )
    metadata["log_path"] = str(log_path)
    return metadata


def _prepare_planned_slides(slide_plan: dict[str, Any], content_analysis: dict[str, Any]) -> list[dict[str, Any]]:
    threshold = float(os.getenv("SLIDE_TEMPLATE_FIT_THRESHOLD", str(DEFAULT_TEMPLATE_FIT_THRESHOLD)))
    slides: list[dict[str, Any]] = []
    for index, item in enumerate(slide_plan.get("slides") or [], start=1):
        planned = dict(item)
        planned["slide_number"] = planned.get("slide_number") or index
        matches = find_relevant_patterns_for_planned_slide(
            planned,
            content_analysis=content_analysis,
            user_goal=slide_plan.get("user_goal"),
            limit=4,
        )
        planned["candidate_templates"] = matches
        best = matches[0] if matches else {}
        planned["matched_template"] = best
        planned["matching_method"] = best.get("matching_method", "none") if best else "none"
        fit_score = float(best.get("template_fit_score") or best.get("similarity_score") or best.get("match_score") or 0)
        type_aligned = bool(best and best.get("slide_type") == planned.get("slide_type"))
        can_use_template = bool(planned.get("use_template_if_relevant", True))
        planned["generation_strategy"] = (
            "template_shell"
            if can_use_template and best and fit_score >= threshold and (type_aligned or fit_score >= threshold + 0.08)
            else "custom_layout"
        )
        if not planned.get("source_evidence") and planned.get("key_message"):
            planned["source_evidence"] = [planned["key_message"]]
        slides.append(planned)
    return slides


def _speaker_notes(slides: list[dict[str, Any]]) -> list[dict[str, str]]:
    notes = []
    for slide in slides:
        body = build_slide_notes(
            {
                **slide,
                "suggested_content": [
                    f"{block.get('heading', '')}: {block.get('body', '')}"
                    for block in slide.get("content_blocks", [])
                    if isinstance(block, dict)
                ],
                "data_required": slide.get("data_points", []),
                "missing_data_or_research_gaps": slide.get("missing_data_or_assumptions", []),
            }
        )
        if slide.get("speaker_notes"):
            body += "\n\nContent-first speaker notes:\n" + str(slide["speaker_notes"])
        notes.append(
            {
                "slide_number": str(slide.get("slide_number") or len(notes) + 1),
                "slide_title": str(slide.get("slide_title") or f"Slide {len(notes) + 1}"),
                "notes": body,
            }
        )
    return notes


def _save_source_grounding_markdown(slide_plan: dict[str, Any], slides: list[dict[str, Any]], output_path: Path) -> Path:
    path = OUTPUT_DIRECTORIES["generated_deck_metadata"] / f"{output_path.stem}_source_grounding.md"
    lines = [f"# Source Grounding: {output_path.stem}", ""]
    lines.append(f"Policy: {slide_plan.get('source_grounding_policy') or 'Strict source grounding.'}")
    lines.append("")
    for slide in slides:
        lines.append(f"## Slide {slide.get('slide_number')}: {slide.get('slide_title')}")
        lines.append("")
        evidence = slide.get("source_evidence") or []
        if evidence:
            lines.append("Source evidence:")
            lines.extend(f"- {item}" for item in evidence)
        else:
            lines.append("Source evidence: Not explicit; manual review required.")
        missing = slide.get("missing_data_or_assumptions") or []
        if missing:
            lines.append("")
            lines.append("Missing data / assumptions:")
            lines.extend(f"- {item}" for item in missing)
        lines.append("")
    path.write_text("\n".join(lines).strip() + "\n", encoding="utf-8")
    return path


def _save_enriched_slide_plan(slide_plan: dict[str, Any], slides: list[dict[str, Any]], output_path: Path) -> str:
    path = OUTPUT_DIRECTORIES["generated_deck_metadata"] / f"{output_path.stem}_slide_plan.json"
    payload = {**slide_plan, "slides": slides}
    write_json(path, payload)
    return str(path)


def _reference_item(planned_slide: dict[str, Any], reference: dict[str, Any]) -> dict[str, Any]:
    matched = planned_slide.get("matched_template") or {}
    return {
        "deck_name": reference.get("deck_name") or matched.get("deck_name"),
        "slide_number": reference.get("slide_number") or matched.get("slide_number"),
        "slide_title": reference.get("slide_title") or matched.get("slide_title"),
        "slide_type": matched.get("slide_type"),
        "chart_type": matched.get("chart_type"),
        "similarity_score": matched.get("similarity_score"),
        "template_fit_score": matched.get("template_fit_score"),
        "generation_strategy": planned_slide.get("generation_strategy"),
    }


def _dedupe_references(items: list[dict[str, Any]]) -> list[dict[str, Any]]:
    seen = set()
    deduped = []
    for item in items:
        key = (item.get("deck_name"), item.get("slide_number"))
        if key in seen:
            continue
        seen.add(key)
        deduped.append(item)
    return deduped


def _candidate_reference_items(slides: list[dict[str, Any]]) -> list[dict[str, Any]]:
    references = []
    for slide in slides:
        matched = slide.get("matched_template") or {}
        if not matched:
            continue
        references.append(
            {
                "deck_name": matched.get("deck_name"),
                "slide_number": matched.get("slide_number"),
                "slide_title": matched.get("slide_title"),
                "slide_type": matched.get("slide_type"),
                "chart_type": matched.get("chart_type"),
                "similarity_score": matched.get("similarity_score"),
                "template_fit_score": matched.get("template_fit_score"),
                "generation_strategy": slide.get("generation_strategy"),
            }
        )
    return references


def _matching_method(slides: list[dict[str, Any]]) -> str:
    methods = [str(slide.get("matching_method") or "") for slide in slides]
    if any(method == "vector_search" for method in methods):
        return "per_slide_vector_search"
    if any(method == "metadata_match" for method in methods):
        return "per_slide_metadata_match"
    return "custom_layout_only"
