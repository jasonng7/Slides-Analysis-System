from pathlib import Path

from pptx import Presentation


def _shape_text(shape) -> list[str]:
    chunks: list[str] = []

    if hasattr(shape, "text") and shape.text:
        chunks.append(shape.text.strip())

    if getattr(shape, "has_table", False):
        for row in shape.table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                chunks.append(" | ".join(cells))

    if getattr(shape, "shape_type", None) == 6 and hasattr(shape, "shapes"):
        for nested_shape in shape.shapes:
            chunks.extend(_shape_text(nested_shape))

    return [chunk for chunk in chunks if chunk]


def extract_text_from_pptx(pptx_path: Path, output_dir: Path, deck_slug: str) -> list[Path]:
    output_dir.mkdir(parents=True, exist_ok=True)
    presentation = Presentation(pptx_path)
    text_paths: list[Path] = []

    for slide_index, slide in enumerate(presentation.slides, start=1):
        chunks: list[str] = []
        for shape in slide.shapes:
            chunks.extend(_shape_text(shape))

        text = "\n".join(dict.fromkeys(chunks)).strip()
        text_path = output_dir / f"{deck_slug}_slide_{slide_index:03d}.txt"
        text_path.write_text(text + ("\n" if text else ""), encoding="utf-8")
        text_paths.append(text_path)

    return text_paths
