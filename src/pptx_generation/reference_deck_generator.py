from __future__ import annotations

from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

from config import OUTPUT_DIRECTORIES
from pptx_generation.generation_models import GeneratedDeckInput, GeneratedDeckMetadata
from pptx_generation.notes_builder import build_slide_notes
from pptx_generation.placeholder_deck_generator import (
    _deck_title,
    _extract_patterns,
    _pattern_matching_method,
    _referenced_template_slides,
    _user_goal,
    load_recommendation_json,
    normalize_recommendation_to_slide_plans,
)
from pptx_generation.pptx_repository import (
    default_output_path,
    ensure_generated_deck_dirs,
    save_generation_metadata,
    save_speaker_notes_markdown,
)
from pptx_generation.reference_slide_resolver import resolve_reference_slide_image
from pptx_generation.slide_layout_builder import (
    SLIDE_H,
    SLIDE_W,
    add_agenda_slide,
    add_appendix_slide,
    add_cover_slide,
)
from recommendation.recommend_from_content import recommend_from_content


DARK = RGBColor(24, 31, 42)
MID = RGBColor(88, 98, 110)
WHITE = RGBColor(255, 255, 255)
ACCENT = RGBColor(10, 64, 116)
SOFT = RGBColor(245, 247, 250)
PEACH = RGBColor(252, 238, 228)
BORDER = RGBColor(178, 188, 199)


def generate_reference_deck_from_recommendation(
    recommendation_path: str | Path,
    output_path: str | Path | None = None,
) -> dict[str, Any]:
    recommendation = load_recommendation_json(recommendation_path)
    return _generate_reference_deck(
        recommendation=recommendation,
        source_recommendation_path=recommendation.get("_source_recommendation_path") or str(recommendation_path),
        output_path=output_path,
    )


def generate_reference_deck_from_text(
    text: str,
    goal: str | None,
    mode: str = "deck",
    output_path: str | Path | None = None,
) -> dict[str, Any]:
    recommendation = recommend_from_content(text=text, user_goal=goal, mode=mode, prompt_output=False)
    return _generate_reference_deck(
        recommendation=recommendation,
        source_recommendation_path=recommendation.get("recommendation_json_path"),
        output_path=output_path,
    )


def generate_reference_deck_from_file(
    file_path: str | Path,
    goal: str | None,
    mode: str = "deck",
    output_path: str | Path | None = None,
) -> dict[str, Any]:
    recommendation = recommend_from_content(
        file_path=str(file_path),
        user_goal=goal,
        mode=mode,
        prompt_output=False,
    )
    return _generate_reference_deck(
        recommendation=recommendation,
        source_recommendation_path=recommendation.get("recommendation_json_path"),
        output_path=output_path,
    )


def _generate_reference_deck(
    *,
    recommendation: dict[str, Any],
    source_recommendation_path: str | None,
    output_path: str | Path | None,
) -> dict[str, Any]:
    try:
        from pptx import Presentation
    except ImportError as exc:
        raise ImportError(
            "python-pptx is required for reference PPTX generation. "
            "Install it with: .venv/bin/python -m pip install python-pptx"
        ) from exc

    ensure_generated_deck_dirs()
    slide_plans = normalize_recommendation_to_slide_plans(recommendation)
    if not slide_plans:
        raise ValueError("No slide plans found in recommendation JSON.")

    timestamp = datetime.now(timezone.utc).strftime("%Y%m%d_%H%M%S")
    title = _deck_title(recommendation)
    target_path = Path(output_path) if output_path else default_output_path(f"{title}_reference_style", timestamp)
    if not target_path.is_absolute():
        target_path = Path.cwd() / target_path
    target_path.parent.mkdir(parents=True, exist_ok=True)

    deck_input = GeneratedDeckInput(
        title=title,
        user_goal=_user_goal(recommendation),
        mode=recommendation.get("recommended_output_mode", "deck"),
        source_recommendation_path=source_recommendation_path,
        matching_method=recommendation.get("matching_method") or _pattern_matching_method(_extract_patterns(recommendation)),
        slide_patterns=_extract_patterns(recommendation),
        recommendation=recommendation,
    ).as_dict()

    prs = Presentation()
    prs.slide_width = Inches(SLIDE_W)
    prs.slide_height = Inches(SLIDE_H)

    add_cover_slide(prs, {**deck_input, "title": f"{title}\nReference-style draft"})
    add_agenda_slide(prs, slide_plans)

    warnings: list[str] = [
        "Stage 8A uses matched slide PNGs as non-editable backgrounds with editable overlays; actual PPTX slide duplication is not implemented yet."
    ]
    content_slides_with_background = 0
    resolved_references: list[dict[str, Any]] = []

    for index, slide_plan in enumerate(slide_plans, start=1):
        reference = resolve_reference_slide_image(slide_plan.get("matched_template") or {})
        if reference.get("found"):
            content_slides_with_background += 1
            resolved_references.append(
                {
                    "generated_slide_number": index,
                    "deck_name": reference.get("deck_name"),
                    "slide_number": reference.get("slide_number"),
                    "slide_title": reference.get("slide_title"),
                    "slide_image_path": reference.get("slide_image_path"),
                }
            )
        elif reference.get("warning"):
            warnings.append(str(reference["warning"]))
        _add_reference_content_slide(prs, slide_plan, index, reference)

    add_appendix_slide(prs, deck_input, slide_plans)
    prs.save(target_path)

    notes_path = save_speaker_notes_markdown(
        [
            {
                "slide_number": str(index),
                "slide_title": slide_plan.get("slide_title", f"Slide {index}"),
                "notes": build_slide_notes(slide_plan),
            }
            for index, slide_plan in enumerate(slide_plans, start=1)
        ],
        target_path,
    )

    metadata = GeneratedDeckMetadata(
        generated_at=datetime.now(timezone.utc).isoformat(),
        output_path=str(target_path),
        source_recommendation_path=source_recommendation_path,
        user_goal=deck_input.get("user_goal", ""),
        mode=deck_input.get("mode", "deck"),
        number_of_slides=len(prs.slides),
        number_of_content_slides=len(slide_plans),
        matching_method=deck_input.get("matching_method", "metadata_match"),
        referenced_template_slides=_referenced_template_slides(slide_plans),
        warnings=list(dict.fromkeys(warnings)),
        generator_version="stage8a_reference_image_background_v1",
    ).as_dict()
    metadata["speaker_notes_path"] = str(notes_path)
    metadata["reference_background_mode"] = "slide_image_background_with_editable_overlays"
    metadata["content_slides_with_reference_background"] = content_slides_with_background
    metadata["content_slides_without_reference_background"] = len(slide_plans) - content_slides_with_background
    metadata["resolved_reference_backgrounds"] = resolved_references

    metadata_path = save_generation_metadata(metadata, target_path)
    metadata["metadata_path"] = str(metadata_path)

    log_path = OUTPUT_DIRECTORIES["generated_deck_logs"] / f"{target_path.stem}.log"
    log_path.write_text(
        "\n".join(
            [
                f"Generated reference-style deck: {target_path}",
                f"Slides: {len(prs.slides)}",
                f"Reference backgrounds: {content_slides_with_background}/{len(slide_plans)}",
                f"Metadata: {metadata_path}",
            ]
        )
        + "\n",
        encoding="utf-8",
    )
    metadata["log_path"] = str(log_path)
    return metadata


def _add_reference_content_slide(
    prs,
    slide_plan: dict[str, Any],
    slide_number: int,
    reference: dict[str, Any],
) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    image_path = reference.get("slide_image_path")
    if image_path:
        slide.shapes.add_picture(
            str(image_path),
            Inches(0),
            Inches(0),
            width=Inches(SLIDE_W),
            height=Inches(SLIDE_H),
        )
    else:
        _add_fallback_background(slide)

    _add_title_overlay(slide, slide_plan)
    _add_content_overlay(slide, slide_plan)
    _add_reference_footer(slide, slide_plan, reference)


def _add_title_overlay(slide, slide_plan: dict[str, Any]) -> None:
    title = slide_plan.get("slide_title") or "Editable action title"
    box = _add_box(slide, 0.55, 0.23, 11.95, 0.68, title, fill=WHITE, font_size=15, bold=True)
    box.line.color.rgb = WHITE


def _add_content_overlay(slide, slide_plan: dict[str, Any]) -> None:
    content = _overlay_content(slide_plan)
    _add_box(slide, 8.55, 1.08, 3.78, 3.65, content, fill=WHITE, font_size=8)
    notes = build_slide_notes(slide_plan)
    excerpt = notes[:360] + ("..." if len(notes) > 360 else "")
    _add_box(slide, 0.68, 5.87, 11.65, 0.92, f"Editable Build Notes\n{excerpt}", fill=SOFT, font_size=6)


def _add_reference_footer(
    slide,
    slide_plan: dict[str, Any],
    reference: dict[str, Any],
) -> None:
    matched = slide_plan.get("matched_template") or {}
    if reference.get("found"):
        footer = (
            f"Reference background: {reference.get('deck_name')} slide {reference.get('slide_number')} "
            f"| {slide_plan.get('slide_type', 'other')} | {slide_plan.get('matching_method', 'metadata_match')}"
        )
    else:
        footer = (
            f"Missing reference background for {matched.get('deck_name', 'unknown deck')} "
            f"slide {matched.get('slide_number', '?')} | fallback overlay"
        )
    _add_textbox(slide, 0.55, 7.06, 12.1, 0.22, footer, 6, color=MID)


def _overlay_content(slide_plan: dict[str, Any]) -> str:
    lines = [
        "Editable Content Overlay",
        f"Type: {slide_plan.get('slide_type', 'other')}",
        f"Chart: {slide_plan.get('chart_type', 'unknown')}",
    ]
    suggested = slide_plan.get("suggested_content")
    if isinstance(suggested, list) and suggested:
        lines.append("")
        lines.extend(f"- {str(item)}" for item in suggested[:6])
    elif isinstance(suggested, str) and suggested.strip():
        lines.extend(["", suggested])
    else:
        lines.extend(["", "- Replace this panel with source-backed content.", "- Preserve the reference slide's visual rhythm."])
    missing = slide_plan.get("missing_data_or_research_gaps")
    if isinstance(missing, list) and missing:
        lines.extend(["", "Missing / to validate:"])
        lines.extend(f"- {str(item)}" for item in missing[:3])
    return "\n".join(lines)


def _add_fallback_background(slide) -> None:
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = RGBColor(247, 248, 250)
    _add_box(
        slide,
        0.7,
        1.25,
        11.9,
        4.25,
        "Reference slide image not found.\nThis slide uses a clean fallback background.",
        fill=PEACH,
        font_size=14,
        bold=True,
    )


def _add_box(
    slide,
    x: float,
    y: float,
    w: float,
    h: float,
    text: str,
    *,
    fill: RGBColor,
    font_size: int,
    bold: bool = False,
):
    shape = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.RECTANGLE,
        Inches(x),
        Inches(y),
        Inches(w),
        Inches(h),
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = BORDER
    shape.line.width = Pt(0.6)
    _set_text(shape, text, font_size=font_size, bold=bold)
    return shape


def _add_textbox(slide, x: float, y: float, w: float, h: float, text: str, font_size: int, *, color: RGBColor) -> None:
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    _set_text(shape, text, font_size=font_size, color=color)


def _set_text(shape, text: str, *, font_size: int, bold: bool = False, color: RGBColor = DARK) -> None:
    text_frame = shape.text_frame
    text_frame.clear()
    text_frame.word_wrap = True
    text_frame.margin_left = Inches(0.08)
    text_frame.margin_right = Inches(0.08)
    text_frame.margin_top = Inches(0.04)
    text_frame.margin_bottom = Inches(0.04)
    for index, line in enumerate(str(text).splitlines() or [""]):
        paragraph = text_frame.paragraphs[0] if index == 0 else text_frame.add_paragraph()
        paragraph.text = line
        paragraph.font.name = "Aptos"
        paragraph.font.size = Pt(font_size)
        paragraph.font.bold = bold if index == 0 else False
        paragraph.font.color.rgb = color if index == 0 else DARK
        paragraph.alignment = PP_ALIGN.LEFT
