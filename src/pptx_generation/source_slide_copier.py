from __future__ import annotations

from copy import deepcopy
from pathlib import Path
from typing import Any

from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

from config import PROJECT_ROOT
from database.db import get_connection
from database.schema import initialize_database
from pptx_generation.reference_slide_resolver import resolve_reference_slide_image


class SlideCloneError(RuntimeError):
    pass


def resolve_reference_slide_source(matched_template: dict[str, Any]) -> dict[str, Any]:
    deck_name = str(matched_template.get("deck_name") or "").strip()
    slide_number = _safe_int(matched_template.get("slide_number"))
    if not deck_name or slide_number is None:
        return {
            "found": False,
            "warning": "Matched template is missing deck_name or slide_number.",
        }

    initialize_database()
    with get_connection() as connection:
        row = connection.execute(
            """
            SELECT
              d.deck_name,
              d.source_file,
              d.source_type,
              s.slide_number,
              s.slide_title,
              s.slide_image_path
            FROM slides s
            JOIN decks d ON d.id = s.deck_id
            WHERE LOWER(d.deck_name) = LOWER(?) AND s.slide_number = ?
            LIMIT 1
            """,
            (deck_name, slide_number),
        ).fetchone()

    if not row:
        return {
            "found": False,
            "warning": f"Could not find source slide for {deck_name} slide {slide_number}.",
        }

    item = dict(row)
    source_path = _resolve_path(item.get("source_file"))
    if not source_path or not source_path.exists() or source_path.suffix.lower() != ".pptx":
        return {
            "found": False,
            **item,
            "warning": f"Source PPTX is unavailable for {deck_name} slide {slide_number}.",
        }

    item["source_file"] = str(source_path)
    item["found"] = True
    return item


def clone_reference_slide_into(
    target_prs,
    slide_plan: dict[str, Any],
    presentation_cache: dict[str, Any] | None = None,
) -> tuple[Any, dict[str, Any]]:
    from pptx import Presentation

    matched = slide_plan.get("matched_template") or {}
    reference = resolve_reference_slide_source(matched)
    if not reference.get("found"):
        raise SlideCloneError(str(reference.get("warning") or "Reference source slide not found."))

    source_file = str(reference["source_file"])
    cache = presentation_cache if presentation_cache is not None else {}
    source_prs = cache.get(source_file)
    if source_prs is None:
        source_prs = Presentation(source_file)
        cache[source_file] = source_prs

    slide_number = int(reference["slide_number"])
    if slide_number < 1 or slide_number > len(source_prs.slides):
        raise SlideCloneError(f"Slide {slide_number} is outside source deck range.")

    source_slide = source_prs.slides[slide_number - 1]
    target_slide = target_prs.slides.add_slide(target_prs.slide_layouts[6])
    _clear_slide(target_slide)
    _clone_slide_xml(source_slide, target_slide)
    _replace_title_text(target_slide, slide_plan)
    _add_editable_content_panel(target_slide, slide_plan)
    _add_clone_footer(target_slide, slide_plan, reference)
    return target_slide, reference


def clone_reference_slide_as_content_shell(
    target_prs,
    slide_plan: dict[str, Any],
    presentation_cache: dict[str, Any] | None = None,
) -> tuple[Any, dict[str, Any]]:
    """Clone a reference slide, remove old text, and add source-grounded content.

    Stage 8B kept most copied body text, which could leave irrelevant template
    content in the generated deck. This variant treats the source slide as a
    visual shell only: clone shapes, clear text/table cells, then place the
    planned slide content back as editable PowerPoint text.
    """
    from pptx import Presentation

    matched = slide_plan.get("matched_template") or {}
    reference = resolve_reference_slide_source(matched)
    if not reference.get("found"):
        raise SlideCloneError(str(reference.get("warning") or "Reference source slide not found."))

    source_file = str(reference["source_file"])
    cache = presentation_cache if presentation_cache is not None else {}
    source_prs = cache.get(source_file)
    if source_prs is None:
        source_prs = Presentation(source_file)
        cache[source_file] = source_prs

    slide_number = int(reference["slide_number"])
    if slide_number < 1 or slide_number > len(source_prs.slides):
        raise SlideCloneError(f"Slide {slide_number} is outside source deck range.")

    source_slide = source_prs.slides[slide_number - 1]
    target_slide = target_prs.slides.add_slide(target_prs.slide_layouts[6])
    _clear_slide(target_slide)
    _clone_slide_xml(source_slide, target_slide)
    _clear_copied_text(target_slide)
    _add_content_first_overlay(target_slide, slide_plan, reference)
    return target_slide, reference


def add_image_fallback_slide(target_prs, slide_plan: dict[str, Any]) -> tuple[Any, dict[str, Any]]:
    from pptx_generation.reference_deck_generator import _add_reference_content_slide

    reference = resolve_reference_slide_image(slide_plan.get("matched_template") or {})
    _add_reference_content_slide(target_prs, slide_plan, int(slide_plan.get("slide_number") or 1), reference)
    return target_prs.slides[-1], reference


def _clone_slide_xml(source_slide, target_slide) -> None:
    rel_map = _copy_relationships(source_slide, target_slide)
    target_tree = target_slide.shapes._spTree

    # Copy the source slide background element when present. This preserves
    # simple theme/background fills while avoiding slide-master mutation.
    source_bg = getattr(source_slide.element.cSld, "bg", None)
    if source_bg is not None:
        target_bg = getattr(target_slide.element.cSld, "bg", None)
        if target_bg is not None:
            target_slide.element.cSld.remove(target_bg)
        target_slide.element.cSld.insert(0, deepcopy(source_bg))

    for shape in source_slide.shapes:
        element = deepcopy(shape.element)
        _remap_relationship_ids(element, rel_map)
        target_tree.insert_element_before(element, "p:extLst")


def _copy_relationships(source_slide, target_slide) -> dict[str, str]:
    rel_map: dict[str, str] = {}
    for old_rid, rel in source_slide.part.rels.items():
        if rel.reltype.endswith("/notesSlide") or rel.reltype.endswith("/slideLayout"):
            continue
        try:
            target = rel.target_ref if rel.is_external else rel.target_part
            new_rid = target_slide.part.relate_to(target, rel.reltype, rel.is_external)
            rel_map[old_rid] = new_rid
        except Exception as exc:
            raise SlideCloneError(f"Could not copy slide relationship {old_rid}: {exc}") from exc
    return rel_map


def _remap_relationship_ids(element, rel_map: dict[str, str]) -> None:
    if not rel_map:
        return
    for node in element.iter():
        for attr_name, attr_value in list(node.attrib.items()):
            if attr_value in rel_map:
                node.attrib[attr_name] = rel_map[attr_value]


def _replace_title_text(slide, slide_plan: dict[str, Any]) -> None:
    title = str(slide_plan.get("slide_title") or "").strip()
    if not title:
        return
    text_shapes = [
        shape
        for shape in slide.shapes
        if getattr(shape, "has_text_frame", False) and (shape.text or "").strip()
    ]
    if not text_shapes:
        _add_textbox(slide, 0.55, 0.25, 11.8, 0.65, title, 15, bold=True)
        return

    def score(shape) -> tuple[int, int, int]:
        top_score = 1 if shape.top < Inches(1.6) else 0
        width_score = 1 if shape.width > Inches(5.0) else 0
        area = int(shape.width * shape.height)
        return (top_score, width_score, area)

    title_shape = sorted(text_shapes, key=score, reverse=True)[0]
    text_frame = title_shape.text_frame
    _set_text_frame(text_frame, title, font_size=15, bold=True)


def _add_editable_content_panel(slide, slide_plan: dict[str, Any]) -> None:
    content = _panel_content(slide_plan)
    _add_box(slide, 8.62, 1.05, 3.62, 3.45, content, font_size=7, fill=RGBColor(255, 255, 255))
    notes = str(slide_plan.get("slide_objective") or slide_plan.get("storyline") or "")
    if notes:
        _add_box(
            slide,
            0.68,
            5.98,
            11.58,
            0.72,
            f"Editable build note: {notes[:300]}",
            font_size=6,
            fill=RGBColor(246, 248, 251),
        )


def _add_clone_footer(slide, slide_plan: dict[str, Any], reference: dict[str, Any]) -> None:
    footer = (
        f"Editable cloned reference: {reference.get('deck_name')} slide {reference.get('slide_number')} "
        f"| {slide_plan.get('slide_type', 'other')} | {slide_plan.get('matching_method', 'metadata_match')}"
    )
    _add_textbox(slide, 0.55, 7.05, 12.1, 0.22, footer, 6, color=RGBColor(88, 98, 110))


def _panel_content(slide_plan: dict[str, Any]) -> str:
    lines = [
        "Editable replacement content",
        f"Type: {slide_plan.get('slide_type', 'other')}",
        f"Chart: {slide_plan.get('chart_type', 'unknown')}",
    ]
    suggested = slide_plan.get("suggested_content")
    if isinstance(suggested, list) and suggested:
        lines.extend(["", *[f"- {item}" for item in suggested[:6]]])
    elif isinstance(suggested, str) and suggested.strip():
        lines.extend(["", suggested])
    else:
        lines.extend(["", "- Replace copied text/visuals with source-backed content."])
    return "\n".join(str(line) for line in lines)


def _clear_copied_text(slide) -> None:
    for shape in slide.shapes:
        if getattr(shape, "has_text_frame", False):
            shape.text_frame.clear()
        if getattr(shape, "has_table", False):
            for row in shape.table.rows:
                for cell in row.cells:
                    cell.text = ""


def _add_content_first_overlay(slide, slide_plan: dict[str, Any], reference: dict[str, Any]) -> None:
    title = str(slide_plan.get("slide_title") or "Source-grounded slide").strip()
    _add_textbox(slide, 0.55, 0.28, 12.15, 0.58, title, 16, bold=True)
    key_message = str(slide_plan.get("key_message") or slide_plan.get("slide_objective") or "").strip()
    if key_message:
        _add_box(slide, 0.65, 1.05, 12.0, 0.7, key_message, font_size=9, fill=RGBColor(236, 244, 251))

    blocks = _overlay_blocks(slide_plan)
    if _is_matrix_like(slide_plan):
        headers = ["Area", "Source-backed content", "Implication / gap"]
        rows = []
        for block in blocks[:5]:
            rows.append([block["heading"], block["body"], block["evidence"]])
        _add_table(slide, 0.72, 2.0, 11.85, 3.65, headers, rows)
    else:
        for index, block in enumerate(blocks[:4]):
            x = 0.72 + (index % 2) * 6.05
            y = 2.0 + (index // 2) * 1.55
            _add_box(
                slide,
                x,
                y,
                5.62,
                1.25,
                f"{block['heading']}\n{block['body']}",
                font_size=8,
                fill=RGBColor(255, 255, 255),
            )

    evidence = _string_items(slide_plan.get("source_evidence"))[:2]
    missing = _string_items(slide_plan.get("missing_data_or_assumptions"))[:1]
    bottom = "Source evidence: " + ("; ".join(evidence) if evidence else "source basis to confirm")
    if missing:
        bottom += "\nMissing / assumption: " + "; ".join(missing)
    _add_box(slide, 0.65, 6.0, 12.0, 0.72, bottom, font_size=6, fill=RGBColor(248, 250, 252))
    _add_clone_footer(slide, slide_plan, reference)


def _overlay_blocks(slide_plan: dict[str, Any]) -> list[dict[str, str]]:
    raw = slide_plan.get("content_blocks")
    blocks: list[dict[str, str]] = []
    if isinstance(raw, list):
        for item in raw:
            if isinstance(item, dict):
                blocks.append(
                    {
                        "heading": str(item.get("heading") or item.get("title") or "Finding"),
                        "body": str(item.get("body") or item.get("content") or ""),
                        "evidence": str(item.get("evidence") or ""),
                    }
                )
            elif str(item).strip():
                blocks.append({"heading": str(item), "body": "", "evidence": ""})
    while len(blocks) < 3:
        evidence = _string_items(slide_plan.get("source_evidence"))
        idx = len(blocks)
        blocks.append(
            {
                "heading": f"Source-backed point {idx + 1}",
                "body": evidence[idx] if idx < len(evidence) else str(slide_plan.get("key_message") or ""),
                "evidence": evidence[idx] if idx < len(evidence) else "",
            }
        )
    return blocks


def _is_matrix_like(slide_plan: dict[str, Any]) -> bool:
    text = " ".join(
        str(slide_plan.get(key) or "").lower()
        for key in ["slide_type", "chart_type", "slide_title", "visual_approach"]
    )
    return any(term in text for term in ["matrix", "table", "benchmark", "comparison", "options", "risk"])


def _add_table(slide, x: float, y: float, w: float, h: float, headers: list[str], rows: list[list[str]]) -> None:
    row_count = max(2, min(len(rows) + 1, 8))
    table_shape = slide.shapes.add_table(row_count, len(headers), Inches(x), Inches(y), Inches(w), Inches(h))
    table = table_shape.table
    for col, header in enumerate(headers):
        cell = table.cell(0, col)
        cell.text = header
        _format_cell(cell, fill=RGBColor(18, 39, 76), color=RGBColor(255, 255, 255), bold=True)
    for row_index, row in enumerate(rows[: row_count - 1], start=1):
        for col in range(len(headers)):
            cell = table.cell(row_index, col)
            cell.text = row[col] if col < len(row) else ""
            _format_cell(cell, fill=RGBColor(255, 255, 255), color=RGBColor(24, 31, 42))


def _format_cell(cell, *, fill: RGBColor, color: RGBColor, bold: bool = False) -> None:
    cell.fill.solid()
    cell.fill.fore_color.rgb = fill
    cell.margin_left = Inches(0.06)
    cell.margin_right = Inches(0.06)
    cell.margin_top = Inches(0.03)
    cell.margin_bottom = Inches(0.03)
    for paragraph in cell.text_frame.paragraphs:
        paragraph.font.name = "Aptos"
        paragraph.font.size = Pt(7.0)
        paragraph.font.bold = bold
        paragraph.font.color.rgb = color


def _string_items(value: object) -> list[str]:
    if isinstance(value, list):
        return [str(item) for item in value if str(item).strip()]
    if isinstance(value, str) and value.strip():
        return [value]
    return []


def _clear_slide(slide) -> None:
    for shape in list(slide.shapes):
        shape.element.getparent().remove(shape.element)


def _add_box(slide, x: float, y: float, w: float, h: float, text: str, *, font_size: int, fill: RGBColor):
    shape = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.RECTANGLE,
        Inches(x),
        Inches(y),
        Inches(w),
        Inches(h),
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = RGBColor(178, 188, 199)
    shape.line.width = Pt(0.5)
    _set_text_frame(shape.text_frame, text, font_size=font_size)
    return shape


def _add_textbox(
    slide,
    x: float,
    y: float,
    w: float,
    h: float,
    text: str,
    font_size: int,
    *,
    color: RGBColor = RGBColor(24, 31, 42),
    bold: bool = False,
) -> None:
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    _set_text_frame(shape.text_frame, text, font_size=font_size, color=color, bold=bold)


def _set_text_frame(
    text_frame,
    text: str,
    *,
    font_size: int,
    color: RGBColor = RGBColor(24, 31, 42),
    bold: bool = False,
) -> None:
    text_frame.clear()
    text_frame.word_wrap = True
    text_frame.margin_left = Inches(0.08)
    text_frame.margin_right = Inches(0.08)
    text_frame.margin_top = Inches(0.04)
    text_frame.margin_bottom = Inches(0.04)
    for index, line in enumerate(str(text).splitlines() or [""]):
        paragraph = text_frame.paragraphs[0] if index == 0 else text_frame.add_paragraph()
        paragraph.text = line
        paragraph.font.name = "Aptos"
        paragraph.font.size = Pt(font_size)
        paragraph.font.bold = bold if index == 0 else False
        paragraph.font.color.rgb = color
        paragraph.alignment = PP_ALIGN.LEFT


def _resolve_path(value: object) -> Path | None:
    if not value:
        return None
    path = Path(str(value)).expanduser()
    if not path.is_absolute():
        path = PROJECT_ROOT / path
    return path.resolve()


def _safe_int(value: object) -> int | None:
    try:
        return int(value)
    except (TypeError, ValueError):
        return None
